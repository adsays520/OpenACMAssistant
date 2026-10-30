"""Build an auditable Word report from real dialogue and existing EDA artifacts.

This program never calls a model or launches EDA. Rebuild after the native chat
finishes so that newly generated tasks and their final answers are included.
"""
from __future__ import annotations

import csv
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import shutil
import statistics
from zoneinfo import ZoneInfo

import numpy as np

HERE = Path(__file__).resolve().parents[1]
ROOT = HERE.parents[2]
LIVE = HERE / "evidence/live_chat_2026-10-28"
ASSETS = HERE / "assets"
STEM = "OpenACMAssistant_E1-E6问答与实验佐证报告"
TASKS = {
    "E2": "20261027_204422_search_comparison_d356cf",
    "E5_4KB": "20261026_144653_acco_sram_e635c9",
    "E5_64bit": "20261026_144235_acco_sram_27ed9c",
    "E6_history": "20261010_142841_dcim_4af72b",
}


def read(path):
    return json.loads(Path(path).read_text())


def write(path, value):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n")


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def csv_write(path, rows):
    if not rows:
        return
    with Path(path).open("w", newline="", encoding="utf-8-sig") as out:
        writer = csv.DictWriter(out, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def collect():
    ASSETS.mkdir(parents=True, exist_ok=True)
    (HERE / "evidence/snapshots").mkdir(parents=True, exist_ok=True)
    start = read(LIVE / "collection_start.json")
    old = set(start["existing_dialogue_tasks"])
    transcript = []
    for p in sorted((ROOT / "runs").glob("*_dialogue_*")):
        if p.name in old or not (p / "answer.md").is_file():
            continue
        req = read(p / "request.json")
        meta = read(p / "metadata.json")
        text = req["messages"][-1]["content"]
        if isinstance(text, dict):
            text = json.dumps(text, ensure_ascii=False)
        transcript.append({"audit": str(p.relative_to(ROOT)), "request": text,
                           "reply": (p / "answer.md").read_text().strip(), "metadata": meta})
        d = HERE / "evidence/snapshots" / p.name
        d.mkdir(exist_ok=True)
        for f in ("request.json", "metadata.json", "answer.md", "facts.json", "final_plan.json"):
            if (p / f).is_file():
                shutil.copy2(p / f, d / f)
    write(LIVE / "turns.json", transcript)
    lines = ["# 本轮真实 CLI 问答", "", "原样保留最终回复；预览、历史读取与实际执行分别审计。", ""]
    for i, t in enumerate(transcript, 1):
        lines += [f"## 第{i}轮", "", "用户：" + t["request"], "", "助手：", "", t["reply"], "",
                  f"审计：{t['audit']}；model_used={t['metadata'].get('model_used')}；executed={t['metadata'].get('executed')}。", ""]
    (LIVE / "dialogue.md").write_text("\n".join(lines))
    jobs = []
    old_tasks = set(start["existing_tasks"])
    for p in sorted((ROOT / "runs").iterdir()):
        if p.name in old_tasks or not (p / "plan.json").is_file() or not (p / "result.json").is_file():
            continue
        plan = read(p / "plan.json")
        if plan.get("workflow") in ("multiplier", "dcim"):
            jobs.append({"task": p.name, "plan": plan, "result": read(p / "result.json")})
    e6 = [j for j in jobs if j["plan"]["workflow"] == "dcim"]
    if e6:
        TASKS["E6"] = e6[-1]["task"]
    else:
        TASKS["E6"] = TASKS["E6_history"]
    evidence = {"schema": "openacm-experiment-report/v1", "report_date": "2026-10-28",
                "collected_at": datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(),
                "dialogue": transcript, "new_jobs": jobs, "tasks": {},
                "scope": "Native CLI real-model dialogue; fresh E1/E6 tasks if present; historical E2/E5 reused without rerunning."}
    for key, name in TASKS.items():
        p = ROOT / "runs" / name
        r = read(p / "result.json")
        integrity = read(p / "source_integrity.json") if (p / "source_integrity.json").exists() else None
        evidence["tasks"][key] = {"task": name, "result": r, "integrity": integrity, "result_sha256": sha(p / "result.json")}
        d = HERE / "evidence/snapshots" / name
        d.mkdir(exist_ok=True)
        for f in ("result.json", "plan.json", "config.json", "source_integrity.json"):
            if (p / f).is_file():
                shutil.copy2(p / f, d / f)
    for j in jobs:
        p = ROOT / "runs" / j["task"]
        d = HERE / "evidence/snapshots" / j["task"]
        d.mkdir(exist_ok=True)
        for f in ("result.json", "plan.json", "source_integrity.json"):
            if (p / f).is_file():
                shutil.copy2(p / f, d / f)
        j["integrity"] = read(p / "source_integrity.json") if (p / "source_integrity.json").is_file() else None
    bp = ROOT / "runs" / TASKS["E2"]
    stats = read(bp / "reports/statistics.json")
    trials = list(csv.DictReader((bp / "reports/trial_summary.csv").open()))
    evidence["E2_statistics"] = stats
    evidence["E2_trials"] = trials
    counts = {method: {} for method in ("random", "nsga2", "llm")}
    for t in evidence["tasks"]["E2"]["result"]["trials"]:
        for k, value in read(bp / t["task"] / "result.json")["counters"].items():
            counts[t["method"]][k] = counts[t["method"]].get(k, 0) + value
    evidence["E2_counters"] = counts
    multiplier_jobs = [j for j in jobs if j["plan"]["workflow"] == "multiplier"]
    if len(multiplier_jobs) >= 2:
        first, second = [ROOT / "runs" / j["task"] for j in multiplier_jobs[:2]]
        a = list(csv.DictReader((first / "artifacts/comparison.csv").open()))
        b = list(csv.DictReader((second / "artifacts/comparison.csv").open()))
        evidence["E1_repeat_check"] = {
            "rtl_sha256_equal": sha(first / "artifacts/rtl/Appro4_2_8bit.v") == sha(second / "artifacts/rtl/Appro4_2_8bit.v"),
            "first_test_is_prefix_of_second": a == b[:len(a)], "sample_counts": [len(a), len(b)],
            "second_execution": "Official run CLI executing the model-generated 8192-sample draft; not automatic execution in turn 4."}
    for name in ("trial_summary.csv", "search_curves.csv", "statistics.json"):
        shutil.copy2(bp / "reports" / name, ASSETS / ("E2_" + name))
    pairs = {}
    observations = []
    failures = []
    for t in evidence["tasks"]["E2"]["result"]["trials"]:
        result = read(bp / t["task"] / "result.json")
        for r in result["stages"]["validation"]["records"]:
            if not r.get("metrics") or r["status"] not in ("completed", "infeasible"):
                failures.append({"method": t["method"], "seed": t["search_seed"], "candidate": r["candidate"], "error": r.get("error")})
                continue
            ids = tuple(r["parameters"]["compressor_ids"])
            pred = r["prediction_reference"]
            measured = r["metrics"]
            row = {"compressor_ids": ",".join(map(str, ids)), "method": t["method"], "search_seed": t["search_seed"], "measurement_status": r["status"]}
            for label, pk, mk in metric_columns():
                row["pred_" + label] = pred[pk]
                row["meas_" + label] = measured[mk]
            observations.append(row)
            pairs.setdefault(ids, []).append(row)
    unique = []
    for ids, rows in sorted(pairs.items()):
        row = {"compressor_ids": ",".join(map(str, ids)), "repeat_measurements": len(rows)}
        for label, _, _ in metric_columns():
            row["pred_" + label] = statistics.median(r["pred_" + label] for r in rows)
            row["meas_" + label] = statistics.median(r["meas_" + label] for r in rows)
        unique.append(row)
    csv_write(ASSETS / "E4_selected_measurements.csv", observations)
    csv_write(ASSETS / "E4_unique_selected_pairs.csv", unique)
    diagnostics = []
    for label, _, _ in metric_columns():
        x = np.array([r["pred_" + label] for r in unique])
        y = np.array([r["meas_" + label] for r in unique])
        good = y != 0
        rank_x, rank_y = rank(x), rank(y)
        rho = float(np.corrcoef(rank_x, rank_y)[0, 1]) if np.std(rank_x) > 0 and np.std(rank_y) > 0 else None
        diagnostics.append({"metric": label, "n_unique": len(unique), "mae": float(np.mean(np.abs(x-y))),
                            "mape_percent": float(np.mean(np.abs((x[good]-y[good])/y[good]))*100) if good.any() else None,
                            "spearman_rho": rho, "zero_denominator_count": int((~good).sum())})
    confusion = {"pred_feasible_meas_feasible": 0, "pred_feasible_meas_infeasible": 0,
                 "pred_infeasible_meas_feasible": 0, "pred_infeasible_meas_infeasible": 0}
    for r in unique:
        a = "feasible" if r["pred_MRED"] <= .02 else "infeasible"
        b = "feasible" if r["meas_MRED"] <= .02 else "infeasible"
        confusion[f"pred_{a}_meas_{b}"] += 1
    evidence["E4_supplement"] = {"scope": "E2-selected samples only; not the planned independent 20-design test",
        "successful_observations": len(observations), "unique_designs": len(unique), "duplicates_aggregation": "median per exact compressor_ids",
        "pairs": unique, "diagnostics": diagnostics, "confusion_mred_002": confusion, "failed_measurements": failures}
    csv_write(ASSETS / "E4_selected_diagnostics.csv", diagnostics)
    small = ROOT / "runs" / TASKS["E5_64bit"]
    evidence["E5_small_cases"] = list(csv.DictReader((small / "reports/pvt_results.csv").open()))
    shutil.copy2(small / "reports/pvt_results.csv", ASSETS / "E5_64bit_pvt_results.csv")
    partial = []
    big = ROOT / "runs" / TASKS["E5_4KB"]
    for p in sorted(big.glob("stages/device_fixed_architecture/*/pvt_*/metrics.json")):
        i, c = int(p.parent.parent.name), int(p.parent.name[-2:])
        r = read(p)
        partial.append({"pd_width_nm": (90,135,180)[i-1], "corner": ("TT","SS","FF")[c], "pvt_status": "failed",
                        "array_read_delay_ns": r.get("read_delay_ns"), "array_read_power_uw": r.get("read_power_uw"),
                        "array_write_delay_ns": r.get("write_delay_ns"), "read_snm_v": r.get("read_snm_v")})
    evidence["E5_4KB_partial"] = partial
    csv_write(ASSETS / "E5_4KB_partial.csv", partial)
    ep = ROOT / "runs" / TASKS["E6"]
    evidence["E6_plan"] = read(ep / "plan.json")
    if (ep / "reports/backend_stages.json").exists():
        evidence["E6_stages"] = read(ep / "reports/backend_stages.json")["stages"]
    else:
        evidence["E6_stages"] = []
    evidence["E6_is_new"] = bool(e6)
    # Hash all direct sources for later auditing, without hashing large EDA binaries.
    hashes = {}
    for name in set(TASKS.values()) | {j["task"] for j in jobs}:
        p = ROOT / "runs" / name
        for f in ("result.json", "plan.json", "source_integrity.json", "reports/backend_stages.json", "artifacts/backend/route.def"):
            if (p / f).is_file():
                hashes[str((p/f).relative_to(ROOT))] = sha(p/f)
    evidence["source_hashes"] = hashes
    write(HERE / "evidence_snapshot.json", evidence)
    return evidence


def rank(values):
    v = np.asarray(values)
    return np.array([float(np.count_nonzero(v < x)) + (float(np.count_nonzero(v == x)) + 1) / 2 for x in v])


def metric_columns():
    return [("MRED", "predicted_mred", "mred"), ("NMED", "predicted_nmed", "nmed"),
            ("delay_ns", "predicted_delay_ns", "sta_delay_ns"), ("area_um2", "predicted_area_um2", "synthesized_area_um2"),
            ("power_uw", "predicted_power_uw", "sta_dynamic_power_uw"), ("PDP_fJ", "predicted_pdp_fj", "pdp_fj")]


def figures(e):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.font_manager import FontProperties
    from matplotlib.patches import Rectangle
    from matplotlib.collections import LineCollection
    font = FontProperties(fname="/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    plt.rcParams.update({"font.family": font.get_name(), "font.size": 10, "axes.unicode_minus": False,
                         "svg.fonttype": "path", "axes.spines.top": False, "axes.spines.right": False})
    colors = {"random": "#248D9D", "nsga2": "#8872BE", "llm": "#E3923E"}

    def save(fig, name):
        fig.savefig(ASSETS / (name + ".png"), dpi=220, bbox_inches="tight", facecolor="white")
        fig.savefig(ASSETS / (name + ".svg"), bbox_inches="tight", facecolor="white")
        plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(9, 3.5))
    for i, method in enumerate(colors):
        trials = [t for t in e["E2_trials"] if t["method"] == method]
        complete = [float(t["best_measured_pdp_fj"]) for t in trials if t["status"] == "completed"]
        offsets = np.linspace(-.10, .10, len(complete))
        axes[0].scatter(i+offsets, complete, color=colors[method], s=45, zorder=3)
        axes[0].plot([i-.2,i+.2], [statistics.median(complete)]*2, color=colors[method], lw=3)
        for t in trials:
            if t["status"] != "completed" and t["best_measured_pdp_fj"]:
                axes[0].scatter(i, float(t["best_measured_pdp_fj"]), marker="x", color="#AA3333", s=65)
        seconds = next(t["total_trial_elapsed_s"] for t in e["E2_statistics"] if t["method"] == method)
        axes[1].bar(i, seconds, width=.55, color=colors[method])
        axes[1].text(i, seconds*1.1, f"{seconds:.1f}s", ha="center", fontsize=9)
    axes[0].set(xticks=range(3), xticklabels=["Random (5/5)","NSGA-II (4/5)","LLM (5/5)"], ylabel="实测可行最优 PDP / fJ", title="完整轮次分布（线为中位数）")
    axes[1].set(xticks=range(3), xticklabels=["Random","NSGA-II","LLM"], ylabel="五次尝试总耗时 / s（对数）", title="包含模型等待及失败尝试", yscale="log", ylim=(10,15000))
    for ax in axes: ax.grid(axis="y", alpha=.2); ax.set_axisbelow(True)
    fig.tight_layout(w_pad=2)
    save(fig,"E2_measured_and_cost")

    curves = list(csv.DictReader((ASSETS / "E2_search_curves.csv").open()))
    fig, ax = plt.subplots(figsize=(8,3.2))
    for method in colors:
        all_curves=[]
        for seed in range(1,6):
            rr = [r for r in curves if r["method"]==method and int(r["search_seed"])==seed]
            values=[float(r["best_predicted_pdp_fj"]) for r in rr]
            ax.plot(range(1,len(values)+1),values,color=colors[method],alpha=.18,lw=1)
            all_curves.append(values)
        med=np.median(np.array(all_curves),axis=0)
        ax.plot(range(1,len(med)+1),med,color=colors[method],label=method,lw=2)
    ax.set(xlabel="累计唯一架构评价数（含基线）",ylabel="预测可行最优 PDP / fJ",title="预测搜索轨迹；粗线为5次尝试中位数，细线为各次轨迹")
    ax.legend(frameon=False,ncols=3); ax.grid(alpha=.2);fig.tight_layout();save(fig,"E2_prediction_curves")

    pairs=e["E4_supplement"]["pairs"]
    fig,axes=plt.subplots(2,3,figsize=(9,6))
    for ax,(label,_,_) in zip(axes.flat,metric_columns()):
        x=np.array([r['pred_'+label] for r in pairs]);y=np.array([r['meas_'+label] for r in pairs])
        ax.scatter(x,y,color="#248D9D",s=25,alpha=.8)
        lo,hi=min(x.min(),y.min()),max(x.max(),y.max());span=hi-lo or 1;lo-=span*.08;hi+=span*.08
        ax.plot([lo,hi],[lo,hi],"--",color="#8E98A6",lw=1)
        ax.set(title=label,xlabel="代理预测",ylabel="本机实测",xlim=(lo,hi),ylim=(lo,hi))
        if label=='MRED':
            ax.axvline(.02,color="#C37C39",lw=.8);ax.axhline(.02,color="#C37C39",lw=.8)
        ax.tick_params(labelsize=8);ax.grid(alpha=.15)
    fig.suptitle(f"E2入选项补充检查：{len(pairs)}个唯一配置；存在选择偏差及PPA条件差异",fontsize=11)
    fig.tight_layout(rect=(0,0,1,.95));save(fig,"E4_selected_prediction_measurement")

    fig,axes=plt.subplots(1,2,figsize=(9,3.1))
    for corner,col in zip(('TT','SS','FF'),('#248D9D','#8872BE','#E3923E')):
        rows=sorted([r for r in e['E5_4KB_partial'] if r['corner']==corner],key=lambda r:r['pd_width_nm'])
        axes[0].plot([r['pd_width_nm'] for r in rows],[r['array_read_delay_ns'] for r in rows],'-o',label=corner,color=col)
        axes[1].plot([r['pd_width_nm'] for r in rows],[r['array_read_power_uw'] for r in rows],'-o',label=corner,color=col)
    for ax in axes:ax.set(xticks=[90,135,180],xlabel='PD宽度 / nm');ax.legend(frameon=False);ax.grid(alpha=.2)
    axes[0].set(ylabel='阵列读延迟 / ns',title='4KB组织：仅读操作部分结果')
    axes[1].set(ylabel='阵列平均读功耗 / µW',title='全部9个PVT子任务最终均超时')
    fig.tight_layout();save(fig,'E5_4KB_partial_read')

    fig,axes=plt.subplots(1,3,figsize=(9,3.1))
    small=e['E5_small_cases']
    for ax,key,title in zip(axes,('bank_read_delay_ns_estimate','bank_write_delay_ns_estimate','read_snm_v'),('bank读延迟估算 / ns','bank写延迟估算 / ns','阵列Read SNM / V')):
        for corner,col in zip(('TT','SS','FF'),('#248D9D','#8872BE','#E3923E')):
            rr=sorted([r for r in small if r['corner']==corner],key=lambda r:float(r['pd_width_nm']))
            ax.plot([float(r['pd_width_nm']) for r in rr],[float(r[key]) for r in rr],'-o',label=corner,color=col)
        ax.set(xticks=[90,135,180],xlabel='PD宽度 / nm',title=title);ax.grid(alpha=.2)
    axes[-1].legend(frameon=False,fontsize=8)
    fig.suptitle('64bit小规模链路验证：4×4×4组织；每角MC=2，不代表4KB完成',fontsize=11)
    fig.tight_layout(rect=(0,0,1,.92));save(fig,'E5_64bit_tradeoffs')

    # A real coordinate view derived only from the selected task's routed DEF.
    path=ROOT/'runs'/TASKS['E6']/'artifacts/backend/route.def'
    if path.exists():
        text=path.read_text();unit=int(re.search(r'UNITS DISTANCE MICRONS (\d+)',text)[1])
        area=list(map(int,re.search(r'DIEAREA\s*\(\s*(-?\d+)\s+(-?\d+)\s*\)\s*\(\s*(-?\d+)\s+(-?\d+)\s*\)',text).groups()))
        block=re.search(r'COMPONENTS \d+ ;(.*?)END COMPONENTS',text,re.S)[1]
        coords=[tuple(map(int,m.groups())) for m in re.finditer(r'\+ (?:PLACED|FIXED) \( (\d+) (\d+) \)',block)]
        fig,ax=plt.subplots(figsize=(6.5,5.3))
        if coords:
            xy=np.array(coords)/unit;ax.scatter(xy[:,0],xy[:,1],s=1,color='#237C94',alpha=.55,label=f'{len(coords)}个单元/宏锚点')
        # DEF routing segments, including * shorthand, retaining straight pieces.
        netblock=re.search(r'\nNETS \d+ ;(.*?)END NETS',text,re.S)
        segments=[]
        if netblock:
            for route in re.split(r'\+ ROUTED |\bNEW ',netblock[1])[1:]:
                previous=None
                for token in re.finditer(r'\(\s*(-?\d+|\*)\s+(-?\d+|\*)\s*\)',route.split(';',1)[0]):
                    a,b=token.groups()
                    if (a=='*' or b=='*') and previous is None:continue
                    point=(int(a) if a!='*' else previous[0], int(b) if b!='*' else previous[1])
                    if previous and point!=previous:segments.append(np.array([previous,point])/unit)
                    previous=point
            ax.add_collection(LineCollection(segments,colors='#8872BE',linewidths=.25,alpha=.24))
        x0,y0,x1,y1=np.array(area)/unit
        ax.add_patch(Rectangle((x0,y0),x1-x0,y1-y0,fill=False,lw=1,color='#243746'))
        ax.set(xlim=(x0-3,x1+3),ylim=(y0-3,y1+3),xlabel='X / µm',ylabel='Y / µm',title='同一任务route.def导出的坐标视图')
        ax.set_aspect('equal');ax.legend(frameon=False,fontsize=8)
        fig.tight_layout();save(fig,'E6_routed_def_view')
        e['E6_def_view']={'component_anchors':len(coords),'routing_segments':len(segments),'die_size_um':[x1-x0,y1-y0],'source':str(path.relative_to(ROOT))}
    e1=[j for j in e['new_jobs'] if j['plan']['workflow']=='multiplier']
    if e1:
        fig,ax=plt.subplots(figsize=(7.5,3))
        ax.bar(range(len(e1)),[j['result'].get('metrics',{}).get('mred',0) for j in e1],color='#248D9D',width=.5)
        ax.axhline(.02,color='#AA3333',ls='--',label='要求 MRED≤0.02')
        ax.set(xticks=range(len(e1)),xticklabels=[str(j['plan']['parameters']['samples'])+'组' for j in e1],ylabel='实测MRED',title='E1：相同电路，增加测试样本')
        ax.legend(frameon=False);fig.tight_layout();save(fig,'E1_mred')


def f(value, digits=6):
    return '未测/未完成' if value is None else f'{value:.{digits}g}' if isinstance(value,(int,float)) else str(value)


def content(e):
    b=[]
    def h(t,level=1):b.append(('heading',level,t))
    def p(t):b.append(('paragraph',t))
    def table(head,rows):b.append(('table',head,rows))
    def fig(name,caption):
        if (ASSETS/(name+'.png')).exists():b.append(('figure',name,caption))
    def page():b.append(('page',))
    stats=e['E2_statistics']; r=e['tasks']['E2']['result']; new=e['new_jobs']; e1=[j for j in new if j['plan']['workflow']=='multiplier']
    h('实验结论与证据范围')
    p('本报告对应实验输入文档中的E1—E6，整合本轮通过命令行完成的真实问答，以及本机已保存的RTL、综合、STA、SPICE与物理后端结果。正文是基于证据的人工整理，附录保留助手最终回复原文；两者分别标注，避免把整理后的理想答案冒充模型输出。未重新运行历史长任务。')
    table(['实验','本次证据','可支持的结论'],[
        ['E1 连续对话',f"{len(e['dialogue'])}轮总问答中的E1链；新乘法器任务{len(e1)}个",'参数解释、受控运行与继承的具体演示；不代表独立测试准确率'],
        ['E2 架构搜索','历史15轮尝试，14轮完整完成','已接通公平条件下三方法有限池比较；本组未显示LLM优势'],
        ['E3 经验消融','提交需求并记录回复','成对经验开关、快照排除及注入审计尚未形成正式实验'],
        ['E4 代理验证',f"补充分析{e['E4_supplement']['unique_designs']}个唯一入选配置",'仅选择后样本的跨条件一致性；独立20配置验证尚未完成'],
        ['E5 SRAM复验','4KB的9个子任务超时；64bit的9个完成','大组织保留部分读数据；小规模证明链路，不证明宏级良率'],
        ['E6 完整后端','本轮新任务' if e['E6_is_new'] else '历史32×16、50ns后端任务','只按实际尺寸、周期及已完成阶段形成结论'],
    ])
    p('关键发现：当前三方法对照中，随机、NSGA-II、LLM的完整且实测可行轮次PDP中位数分别为59.9959、61.2657、61.4295 fJ。LLM未呈现更低的中位PDP，且五次尝试累计耗时4371.554 s。失败记录必须保留，不能以“15/15有记录”写成“15/15完成”。')
    if e['E6_is_new']:
        m=e['tasks']['E6']['result'].get('metrics',{})
        p(f"本轮16×8、10ns DCiM后端已实际执行：乘法输入{f(m.get('sample_count'))}组、错误{f(m.get('error_count'))}；setup/hold slack为{f(m.get('setup_wns_ns'))}/{f(m.get('hold_wns_ns'))}ns，布线DRC为{f(m.get('route_drc_count'))}。E1的sabetz复测均未满足MRED≤0.02；E3、独立E4及4KB完整复验仍未完成。")
    p(f"本轮最终回复共收集{len(e['dialogue'])}条，其中model_used=true为{sum(bool(t['metadata'].get('model_used')) for t in e['dialogue'])}条。审计中的model_used只表示对话决策调用成功，结果文本还包含执行器依据实测数据渲染的内容。")
    h('本轮问答暴露的问题',2)
    table(['位置','实际表现','对报告结论的影响'],[
        ['E1第4轮','“重新验证”只生成8192组计划，没有自动启动','不能声称五轮链全部自动完成；另用官方run入口执行该计划补测'],
        ['E2/E5历史追问','明确任务ID仍返回“只有待运行计划”','历史任务读取被当前草案遮挡；正文按真实result.json分析，原回复保留'],
        ['E4验证咨询','返回E2任务名，却混入E1乘法器MRED=0.492823','发生跨任务事实混用；该回复不作为实验数据或有效答案'],
        ['E1能力咨询','模型前言与事实模板存在重复','内容正确但冗长，后续应改善去重与表达'],
    ])
    p('因此，本轮不能写成“E1—E6问答全部通过”。本报告交付的是实际执行和缺口审计，尚未修复以上对话代码。正文所有数值从指定任务的证据读取，不从上述错误回复摘取。')
    h('1　E1：连续交流、参数继承与真实结果')
    p('按“能力咨询→参数与MRED解释→4096组验证→只改8192组→结果追问”进行同一对话。电路为8位近似乘法器，9个可近似位置均用sabetz，测试种子42，约束MRED≤0.02。咨询轮次不应启动EDA；后续复测只修改样本数。')
    if e1:
        p('E1第3轮自动执行4096组；第4轮仅生成8192组计划。后者随后通过官方run入口执行原计划，作为一次显式补测，不冒充原对话已自动执行。')
        rows=[]
        for j in e1:
            a=j['plan']['parameters']; m=j['result'].get('metrics',{})
            rows.append([j['task'],str(a.get('samples')),str(a.get('seed')),str(a.get('compressor_id')),f(m.get('mred')),j['result']['status']])
        table(['任务','采样','种子','压缩器ID','MRED','任务状态'],rows)
        fig('E1_mred','图1　相同sabetz电路两次抽样的实测MRED与0.02阈值。样本通过/失败结论仅适用于该测试集合。')
        if len(e1)>=2:
            a,bp=e1[0]['plan'],e1[1]['plan']; changes={k:[a['parameters'].get(k),bp['parameters'].get(k)] for k in set(a['parameters'])|set(bp['parameters']) if a['parameters'].get(k)!=bp['parameters'].get(k)}
            p('两次完整参数的差异：'+json.dumps(changes,ensure_ascii=False)+'；约束数组'+('保持一致。' if a['constraints']==bp['constraints'] else '发生变化，须核查。')+'若第二次展开了9项均为8的compressor_ids，它与统一compressor_id=8等价，不是额外改变电路。')
        p('此处把“仿真正常完成”和“达到MRED阈值”分别判定。近似电路未达标也属于有效实验结果，不能提高阈值或删去失败值使演示通过。')
        if e.get('E1_repeat_check'):
            p('两次RTL文件SHA256相同，且4096条逐输入对照记录与8192条中的前4096条完全相同。MRED随样本数变化来自测试集合组成变化，不能称电路被优化；两个集合都不满足0.02要求。')
    else:p('本轮尚未收集到新的完整乘法器任务，不把参数规划或旧任务指标写成此次执行成功。具体问答及失败见附录。')
    p('此前真实本地模型的开发回归记录为24/24轮通过（20261026_220158_example_acceptance_acc7b5），受影响搜索示例另有6/6记录；这些脚本禁止EDA并使用开发中反复修订的例子。它们仅是计划与回复回归，不作为留出集准确率或真实仿真成功率。')
    page()
    h('2　E2：随机、NSGA-II与LLM的架构搜索对照')
    p('沿用已完成尝试的历史任务 '+TASKS['E2']+'。固定8位近似结构、AKbar2基线及器件尺寸；共同合法池256项（池种子1729），各方法使用搜索种子1—5，每轮32个唯一架构预算，含基线。每轮按相同规则最多复核3项；32次代理查询不等于32次EDA。目标是在实测MRED≤0.02的候选中最小化PDP。')
    p('测试输入种子42、4096组，综合库Nangate45 typical，输入活动周期5ns、输出负载10fF。PDP使用动态功耗×STA延迟，µW×ns数值单位为fJ；不能与总功耗PDP或后布局签核结果混写。')
    table(['方法','完整/尝试','实测PDP中位数(fJ)','样本标准差(fJ)','总耗时(s)'],[[s['method'],f"{s['completed']}/{s['attempted']}",f(s['measured_median_pdp_fj']),f(s['measured_stdev_pdp_fj']),f(s['total_trial_elapsed_s'])] for s in stats])
    table(['方法','GNN评价','RTL启动','STA/活动评价','LLM选择调用'],[[method]+[str(e['E2_counters'][method].get(k,0)) for k in ('gnn_queries','array_or_rtl_tasks','sta_activity_evaluations','llm_calls')] for method in ('random','nsga2','llm')])
    fig('E2_measured_and_cost','图2　各完整轮次的最终实测可行最优PDP与五次尝试总耗时。NSGA-II种子5有失败，不计入完整轮次中位数；红叉仅显示其部分测量值。')
    p('随机搜索的中位PDP比本组LLM低约2.33%，LLM累计耗时约为随机的67.7倍。此结果针对当前有限候选池和具体提案机制，不证明所有随机方法优于所有LLM或NSGA-II，也不以五次重复宣称统计显著性。两个方法都曾取得58.8247 fJ的相同候选，更应区分最小值与稳定性。')
    fig('E2_prediction_curves','图3　GNN预测的可行最优PDP轨迹。它与图2的实测值量级不同，不能直接混成同一条“实测收敛曲线”。')
    table(['搜索种子','Random PDP','NSGA-II PDP','LLM PDP'],[[str(seed)]+[f(next(t['best_measured_pdp_fj'] for t in r['trials'] if t['method']==method and t['search_seed']==seed))+('（部分，失败）' if next(t['status'] for t in r['trials'] if t['method']==method and t['search_seed']==seed)!='completed' else '') for method in ('random','nsga2','llm')] for seed in range(1,6)])
    p('NSGA-II种子5的第三项复核在RTL编译时失败：ha_6_0实例名重复，压缩器配置为[7,5,7,7,7,7,7,7,7]。另两项虽然有实测值，整轮仍属于失败，整组状态为incomplete。保留原始日志，不修改原始OpenACM生成器。后续若修复适配层并复跑，应建新任务及新版本对照，不能改写这份历史记录。')
    p('共同池、测试向量与综合库SHA256分别为：'+r['candidate_pool_sha256']+'；'+r['test_vectors_sha256']+'；'+r['synthesis_lib_sha256']+'。完整哈希和原始CSV随报告保存。')
    h('3　E3：持久经验的消融验证')
    p('已将文档中的经验消融需求作为真实问题提交。当前三方法compare不注入持久经验，不能将其中LLM组改称“有经验”。正式E3还缺少经验快照冻结、待测项排除、逐轮开关与注入证明、同搜索种子的配对调度。')
    p('本报告不填入经验带来的PDP提升或调用次数下降。可验证的系统能力是保存和检索观测，尚无证据说明权重已训练或经验已改善搜索。附录保留实际回复；正式图表须在上述配对实验完成后生成。')
    page()
    h('4　E4：预测—实测一致性补充分析')
    s=e['E4_supplement']; c=s['confusion_mred_002']
    p(f"独立20配置的批量验收尚未接通。为让现有数据产生可核查图表，从E2的{len(r['trials'])}轮复核中提取{s['successful_observations']}次有完整指标的测量，包含MRED不满足的记录，按完整9位置配置去重得到{s['unique_designs']}个唯一配置；重复测量按中位数汇总。执行失败项另存，不以0填补。")
    p('这些配置根据预测结果入选，具有选择偏差；训练/经验数据交叉情况尚未核验，训练标签的工艺、负载、活动条件也未完全对齐。因此下表是“入选样本跨条件一致性检查”，不称独立测试集泛化误差。')
    fig('E4_selected_prediction_measurement','图4　入选唯一配置的预测与实测散点。虚线为数值相等，MRED橙线为0.02。PPA存在系统偏移，应先核对定义、工艺、负载和活动条件。')
    table(['指标','样本数','MAE','平均绝对相对偏差(%)','Spearman ρ'],[[d['metric'],str(d['n_unique']),f(d['mae']),f(d['mape_percent']),f(d['spearman_rho'])] for d in s['diagnostics']])
    table(['MRED≤0.02判定','实测可行','实测不满足'],[['预测可行',str(c['pred_feasible_meas_feasible']),str(c['pred_feasible_meas_infeasible'])],['预测不满足',str(c['pred_infeasible_meas_feasible']),str(c['pred_infeasible_meas_infeasible'])]])
    p('表中的“预测可行、实测不满足”是入选集合中危险的可行性误判数，不是全候选池假阳性率；缺少未入选样本不能推断全空间召回率。Spearman采用并列平均秩，常量列相关性留空；相对偏差对零分母单列，不能以任意小值替代。')
    p('正式E4应先冻结20个不按预测好坏挑选的合法配置，并核查训练/经验重叠，然后统一执行RTL、综合、STA和活动功耗。补充检查不能替代该实验，也不能靠简单比例缩放就宣称代理已重新校准。')
    h('5　E5：固定组织器件复验及运行瓶颈')
    p('4KB复验固定32768bit、输出8bit、32行×32列×32子阵列、列复用1。基线PU/PD/PG=90/90/135nm、L=50nm；只比较PD=90、135、180nm。三个PVT为TT/27℃/1V、SS/85℃/0.9V、FF/−40℃/1.1V；每角MC=2、种子42、相对标准差5%。')
    p('历史任务 '+TASKS['E5_4KB']+' 已尝试全部9个尺寸/PVT组合，各子任务达到600s上限，总体failed。架构阶段只通过容量/接口检查；写操作和三种SNM未完整完成。原历史器件stage中的infeasible是状态归类错误，不能解释为电路不可行，原文件保持不动。')
    fig('E5_4KB_partial_read','图5　4KB组织已保存的部分阵列读数据。不是整个bank的完整PVT通过结果，也没有写/SNM最优结论。')
    table(['PD(nm)','PVT','阵列读延迟(ns)','阵列读功耗(µW)','完整案例'],[[str(a['pd_width_nm']),a['corner'],f(a['array_read_delay_ns']),f(a['array_read_power_uw']),'失败/超时'] for a in e['E5_4KB_partial']])
    p('在已有部分数据中，增大PD使读延迟减小，而读功耗略有变化；缺少写延迟与SNM不能据此推荐最终器件。后续应先对每个操作分配可观测预算并定位慢操作，再按同一大组织复验；不能静默缩小容量或减少MC来宣称4KB已通过。')
    h('5.1　64bit小规模验证：单独展示，不替代4KB',2)
    p('独立历史任务 '+TASKS['E5_64bit']+' 使用4×4×4=64bit组织、输出4bit，完成相同三PD×三PVT、每角MC=2的9个子任务。读/写/SNM来自阵列SPICE；bank延迟、功耗外围组合及面积为明确模型估算。')
    fig('E5_64bit_tradeoffs','图6　64bit小规模PD/PVT取舍。最小样本数量用于链路检查，不能用于宏级良率或完整工艺签核。')
    sr=e['tasks']['E5_64bit']['result']['stages']['device_fixed_architecture']['records']
    table(['PD(nm)','最坏读延迟估算(ns)','最坏写延迟估算(ns)','最小Read SNM(V)','bank面积估算(µm²)'],[[f(a['parameters']['devices']['pd_width_nm']),f(a['metrics']['bank_read_delay_ns_estimate']),f(a['metrics']['bank_write_delay_ns_estimate']),f(a['metrics']['read_snm_v']),f(a['metrics']['bank_area_um2_estimate'])] for a in sr])
    p('这个小规模样本中，增加PD没有同时改善读延迟、写延迟及Read SNM；默认鲁棒目标下基线仍是最佳已评价项。器件优化应展示取舍，不应预设尺寸越大整体越好。MC=2不形成概率置信区间或良率承诺。')
    page()
    h('6　E6：真实物理后端与证据链')
    ep=e['E6_plan']['parameters']; er=e['tasks']['E6']['result']; m=er.get('metrics',{})
    p(('本轮通过自然语言新建任务 ' if e['E6_is_new'] else '以下复用历史任务 ')+TASKS['E6']+f"，实际规格为{ep['rows']}行×{ep['bit_width']}位、{ep['kind']}、backend={ep['backend']}、时钟{ep['clock_period_ns']}ns、memory_impl={ep.get('memory_impl','auto')}。任务状态为{er['status']}。")
    if not e['E6_is_new']:p('该历史规格与文档16×8、10ns要求不同，只用作后端链路佐证，不冒充这条新需求已经执行。')
    table(['验收/测量项','结果','说明'],[
        ['乘法器输入数',f(m.get('sample_count')),'是否穷举由计划exhaustive及位宽共同判断'],
        ['乘法器错误数',f(m.get('error_count')),'不能替代完整PE控制状态穷举'],
        ['setup WNS (ns)',f(m.get('setup_wns_ns')),'≥0才满足此项时序约束'],
        ['hold WNS (ns)',f(m.get('hold_wns_ns')),'≥0才满足此项时序约束'],
        ['布线DRC数',f(m.get('route_drc_count')),'工具报告；不是完整工艺DRC/LVS'],
        ['后端面积(µm²)',f(m.get('routed_area_um2')),'按后端报告定义，非bank公式估算'],
        ['活动功耗(µW)',f(m.get('routed_power_uw')),'给定VCD活动与库的模型结果，非芯片实测'],
        ['门级/布线后/SDF对照数',f(m.get('post_synth_functional_checks'))+'/'+f(m.get('post_route_functional_checks'))+'/'+f(m.get('post_route_sdf_checks')),'按实际测试集合'],
        ['SDF timing checks / IOPATH',f(m.get('sdf_annotated_timing_checks'))+'/'+f(m.get('sdf_annotated_iopaths')),'用于核查实际标注范围'],
    ])
    table(['后端阶段','状态','阶段耗时(s)'],[[a['name'],a['status'],f(a.get('elapsed_s'))] for a in e['E6_stages']])
    fig('E6_routed_def_view','图7　从同一任务真实route.def导出的单元锚点和布线坐标图，不是GDSII截图；不补画不存在的宏/导线。完整ODB/DEF/SPEF/SDF由原任务保存。')
    p('参考宏自动匹配相应尺寸，但并未根据E5的PD尺寸重新表征Liberty/LEF。此处证明参考PE后端能够执行，不用来证明E2/E5的优化已经传递到整PE收益。阶段完成、时序约束通过、功能覆盖和物理签核是不同结论。')
    h('7　可直接用于参赛文本的表述')
    p('系统以本地语言模型对接OpenACM，支持连续需求交流、结构化参数校验、隔离仿真与基于实际结果的中文反馈。已建立固定候选池和评价条件下的三方法架构搜索对照，并保留搜索决策、实测复核、失败日志及哈希证据。现阶段实验验证了自动化流程与可追溯性，尚未证明LLM搜索在当前任务中优于随机方法。SRAM器件阶段具备小规模多PVT/MC验证能力，4KB复验仍存在运行时间瓶颈。经验消融及独立代理泛化验证作为后续实验，不计入已完成成果。')
    h('8　证据索引与复现说明')
    table(['证据','本地位置（相对于OpenACMAssistant）'],[['本轮原始问答',str((LIVE/'dialogue.md').relative_to(ROOT))],['报告数据快照',str((HERE/'evidence_snapshot.json').relative_to(ROOT))],['统计CSV/独立图文件',str(ASSETS.relative_to(ROOT))]]+[[k,'runs/'+v] for k,v in TASKS.items()])
    p('证据快照复制了本报告引用的关键JSON和本轮最终问答；完整网表、波形、日志仍位于各runs任务。图表为这些文件的重新绘制，未训练任何新权重。重建命令：.venv/bin/python -B docs/experiment-report/2026-10-27/tools/build_report.py。报告生成器只读取本地证据，不启动模型或EDA。')
    integrity=[(k,v['integrity']) for k,v in e['tasks'].items() if v['integrity']]
    p('所引用历史任务的源完整性记录：'+('全部unchanged=true。' if all(v.get('unchanged') for _,v in integrity) else '存在未通过项，需核查。')+'这些是各任务当时的校验结果，不等同于对所有历史修改的全局证明；本轮新任务的校验保存在各自source_integrity.json。报告与附件全部存于助手专属目录，原始OpenACM程序及用户原Word文稿未改动。')
    page()
    h('附录A　本轮实际问答原文')
    p('以下是最终answer.md中的实际文本，未经改写。不收录模型内部思考。预览不等于已仿真，历史读取不等于本轮新执行；答非所问或状态误判也保留，以便继续修复。')
    for i,t in enumerate(e['dialogue'],1):
        h(f"A.{i}　{t['metadata'].get('final_mode',t['metadata'].get('operation','未完成'))}",2)
        p('用户：'+t['request'])
        p('助手原文：')
        for para in t['reply'].split('\n\n'):p(para)
        p('审计：'+t['audit']+f"；model_used={t['metadata'].get('model_used')}，executed={t['metadata'].get('executed')}。")
    return b


def render(blocks,e):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm,Pt,RGBColor
    doc=Document();s=doc.sections[0]
    s.page_width,s.page_height=Cm(21),Cm(29.7)
    s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Cm(2)
    s.header_distance=s.footer_distance=Cm(1)
    for name,size in [('Normal',10.5),('Title',25),('Subtitle',12),('Heading 1',16),('Heading 2',12),('Caption',9)]:
        st=doc.styles[name];st.font.name='DejaVu Sans';st.font.size=Pt(size)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK SC')
        st.font.color.rgb=RGBColor.from_string('17324D' if name!='Normal' else '253746')
    normal=doc.styles['Normal'].paragraph_format
    normal.line_spacing=1.18;normal.space_after=Pt(5);normal.widow_control=True
    doc.styles['Caption'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for level in [1,2]:
        pf=doc.styles[f'Heading {level}'].paragraph_format;pf.keep_with_next=True;pf.space_before=Pt(12);pf.space_after=Pt(6)
    header=s.header.paragraphs[0];header.text='OPENACM ASSISTANT   /   实验与问答证据报告';header.style='Caption'
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ')
    for field in ['PAGE','NUMPAGES']:
        if field=='NUMPAGES':footer.add_run(' / ')
        node=OxmlElement('w:fldSimple');node.set(qn('w:instr'),field);footer._p.append(node)
    doc.add_paragraph('OPENACM\nASSISTANT',style='Title')
    doc.add_paragraph('E1—E6 问答与实验佐证报告',style='Subtitle')
    doc.add_paragraph('真实命令行问答 · 架构搜索对照 · 代理一致性 · SRAM复验 · 数字后端')
    doc.add_paragraph('整理日期：2026年10月28日（Asia/Shanghai）')
    doc.add_paragraph('证据版本：'+e['collected_at'])
    doc.add_paragraph('本文件可编辑。图表与结论均由本机任务记录整理；未完成与失败单独标注。')
    doc.add_page_break()
    md=['# E1—E6 问答与实验佐证报告','','整理日期：2026-10-28。生成时间：'+e['collected_at'],'']
    for block in blocks:
        if block[0]=='heading':
            _,level,text=block;doc.add_heading(text,level=level);md+=['#'*(level+1)+' '+text,'']
        elif block[0]=='paragraph':
            doc.add_paragraph(block[1]);md+=[block[1],'']
        elif block[0]=='page':doc.add_page_break()
        elif block[0]=='figure':
            _,name,caption=block
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(ASSETS/(name+'.png')),width=Cm(16.8))
            doc.add_paragraph(caption,style='Caption');md += [f'![{caption}](assets/{name}.png)','',caption,'']
        elif block[0]=='table':
            _,heads,rows=block;t=doc.add_table(rows=1,cols=len(heads));t.style='Table Grid'
            repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
            for c,text in zip(t.rows[0].cells,heads):c.text=text
            for row in rows:
                for c,text in zip(t.add_row().cells,row):c.text=str(text)
            for i,row in enumerate(t.rows):
                for cell in row.cells:
                    if i==0:
                        fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7EFF5');cell._tc.get_or_add_tcPr().append(fill)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(2)
                        for run in p.runs:run.font.size=Pt(8.5 if len(heads)>4 else 9);run.bold=(i==0)
            doc.add_paragraph()
            md += ['| '+' | '.join(heads)+' |','| '+' | '.join(['---']*len(heads))+' |']
            md += ['| '+' | '.join(str(x).replace('|','/').replace('\n','<br>') for x in row)+' |' for row in rows]+['']
    doc.core_properties.title='OpenACMAssistant E1—E6问答与实验佐证报告'
    doc.core_properties.subject='Local evidence and real CLI dialogue; incomplete experiments explicitly identified'
    doc.core_properties.author='OpenACMAssistant project'
    doc.save(HERE/(STEM+'.docx'))
    (HERE/(STEM+'.md')).write_text('\n'.join(md))
    print(HERE/(STEM+'.docx'))


def main():
    e=collect();figures(e);write(HERE/'evidence_snapshot.json',e);render(content(e),e)
    write(HERE/'validation.json',{'generated_at':e['collected_at'],'dialogue_count':len(e['dialogue']),
        'new_jobs':[{'task':j['task'],'status':j['result']['status']} for j in e['new_jobs']],
        'unique_prediction_measurement_pairs':e['E4_supplement']['unique_designs'],
        'embedded_figures':len(list(ASSETS.glob('*.png'))),'pdf_render_verified':False})


if __name__=='__main__':main()

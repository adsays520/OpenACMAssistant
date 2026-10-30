"""Extend the supplied report in place in a COPY, preserving its Word structure.

No EDA is launched and no original OpenACM files are modified.
"""
from collections import Counter
from copy import deepcopy
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import shutil
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

HERE = Path(__file__).resolve().parents[1]
ROOT = HERE.parents[2]
SOURCE = ROOT.parent / "技术报告.docx"
STEM = "技术报告_特色功能完善版"
TASK = "20261011_140055_explore_multiplier_fdabd4"
ACCO = "20261025_132333_acco_multiplier_5f789e"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n")


def evidence():
    p = ROOT / "runs" / TASK
    r = json.loads((p / "result.json").read_text())
    rows = r["candidates"]
    best = next(x for x in rows if x["candidate"] == r["best_feasible_candidate"])
    integrity = json.loads((p / "source_integrity.json").read_text())
    counts = dict(Counter(x["status"] for x in rows))
    proposers = dict(Counter(x["proposed_by"] for x in rows))
    assert len(rows) == 128 and counts == {"completed": 18, "infeasible": 86, "failed": 24}
    assert proposers == {"baseline": 1, "deepseek": 127}
    assert integrity["unchanged"]
    for selection in (p / "decisions").rglob("selection.json"):
        choice = json.loads(selection.read_text())
        if choice.get("proposed_by") == "deepseek":
            assert (selection.parent / "request.json").exists()
            assert (selection.parent / "reply.json").exists()
    area0 = rows[0]["metrics"]["synthesized_area_um2"]
    areab = best["metrics"]["synthesized_area_um2"]
    result = {"task": TASK, "kind": r["kind"], "status": r["status"], "budget": r["budget"],
              "candidate_records": len(rows), "status_counts": counts, "proposed_by_counts": proposers,
              "pareto_candidates": r["pareto_candidates"], "best": best, "baseline": rows[0],
              "area_reduction_percent": 100 * (area0-areab)/area0,
              "elapsed_s": sum(x["elapsed_s"] for x in r["sessions"]),
              "stop_reason": r["stop_reason"], "source_integrity": integrity,
              "result_sha256": sha(p / "result.json"), "config": r["config"]}
    ap = ROOT / "runs" / ACCO
    ar = json.loads((ap / "result.json").read_text())
    result["acco_example"] = {"task": ACCO, "status": ar["status"], "counters": ar["counters"],
                              "result_sha256": sha(ap / "result.json")}
    write_json(HERE / "evidence/feature_evidence.json", result)
    for name in (TASK, ACCO):
        d = HERE / "evidence" / name
        d.mkdir(exist_ok=True)
        for f in ("result.json", "source_integrity.json"):
            shutil.copy2(ROOT / "runs" / name / f, d / f)
    return result, rows


def figures(info, rows):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.font_manager import FontProperties
    from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
    font = FontProperties(fname="/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    plt.rcParams.update({"font.family": font.get_name(), "font.size": 10, "axes.unicode_minus": False,
                         "svg.fonttype": "path", "axes.spines.top": False, "axes.spines.right": False})
    navy, teal, amber = "#17324D", "#198A8A", "#D58A39"
    def save(fig, name):
        for ext in ("png", "svg"):
            fig.savefig(HERE / "assets" / (name + "." + ext), dpi=220, bbox_inches="tight", facecolor="white")
        plt.close(fig)
    fig, ax = plt.subplots(figsize=(9,5))
    ax.set(xlim=(0,1),ylim=(0,1));ax.axis("off")
    def box(x,y,w,h,t,fill="#EAF2F7",color=navy):
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle="round,pad=0.012,rounding_size=0.015",facecolor=fill,edgecolor=color,lw=1.2))
        ax.text(x+w/2,y+h/2,t,ha="center",va="center",color=color,linespacing=1.5)
    def arrow(a,b,dashed=False):
        ax.add_patch(FancyArrowPatch(a,b,arrowstyle="-|>",mutation_scale=14,color="#66798A",lw=1.4,linestyle="--" if dashed else "-"))
    box(.13,.80,.74,.15,"用户描述目标、硬约束、固定条件与预算\n保持连续对话，补齐缺项后形成结构化实验")
    box(.035,.48,.28,.19,"构造合法候选菜单\n固定规格与测试条件\n基线 / 组合 / 邻域",fill="#E5F4F1",color=teal)
    box(.36,.48,.28,.19,"本地LLM选择下一项\n读取已评价历史与失败\n输出菜单编号及理由")
    box(.685,.48,.28,.19,"执行对应评价工具\nGNN预测或真实EDA\n指标、状态与证据分开")
    arrow((.5,.79),(.5,.69));arrow((.32,.58),(.35,.58));arrow((.65,.58),(.675,.58))
    box(.24,.15,.52,.19,"更新可行最优与Pareto候选\n继续下一轮；达到预算后停止并总结\n保存决策、日志、参数及经验引用",fill="#FFF5E8",color="#99602A")
    arrow((.825,.47),(.76,.28));arrow((.24,.25),(.175,.47),True)
    ax.text(.5,.045,"局部路径逐项实测；ACCO按架构预测 → 入选复核 → 固定架构器件评价分层执行",ha="center",fontsize=9,color="#66798A")
    fig.tight_layout();save(fig,"free_exploration_loop")

    fig,axes=plt.subplots(1,2,figsize=(9,3.6))
    current=None;curve=[]
    for r in rows:
        if r['status']=='completed':
            value=r['metrics']['synthesized_area_um2'];current=value if current is None else min(current,value)
        curve.append(current)
    axes[0].step(range(1,len(rows)+1),curve,where='post',color=teal,lw=2)
    axes[0].axhline(info['baseline']['metrics']['synthesized_area_um2'],ls='--',color='#93A1AF',lw=1)
    axes[0].set(xlabel='累计候选记录数',ylabel='可行最优综合面积 / µm²',title='目标反馈与预算内的最优记录',xlim=(1,128))
    axes[0].annotate('第6条：291.802 µm²',xy=(6,291.802),xytext=(30,305),arrowprops={'arrowstyle':'->','color':navy},fontsize=9,color=navy)
    for status,color,label in [('infeasible','#B9C2CC','约束未达标（86）'),('completed',teal,'满足约束（18）')]:
        rr=[r for r in rows if r['status']==status]
        axes[1].scatter([r['metrics']['mred'] for r in rr],[r['metrics']['synthesized_area_um2'] for r in rr],s=22,c=color,label=label,alpha=.8)
    best=info['best']['metrics']
    axes[1].scatter([best['mred']],[best['synthesized_area_um2']],s=120,marker='*',c=amber,zorder=4,label='本次最小可行面积')
    axes[1].axvline(.02,ls='--',color='#B56D36',lw=1)
    axes[1].set(xscale='log',xlabel='实测 MRED（对数）',ylabel='综合面积 / µm²',title='固定测试条件下的误差与面积取舍')
    axes[1].legend(frameon=False,fontsize=8,loc='lower left')
    for ax in axes:ax.grid(alpha=.16);ax.set_axisbelow(True)
    fig.tight_layout(w_pad=1.8);save(fig,'free_exploration_results')


def main():
    for f in ("assets", "evidence"):(HERE/f).mkdir(parents=True,exist_ok=True)
    original_hash=sha(SOURCE)
    backup=HERE/'evidence/技术报告_原稿.docx'
    if not backup.exists():shutil.copy2(SOURCE,backup)
    info,rows=evidence();figures(info,rows)
    doc=Document(SOURCE);p=list(doc.paragraphs);tables=list(doc.tables);changes=[]
    body_pr=deepcopy(p[21]._p.pPr)

    def replace(i,text):
        para=p[i]
        assert not para._p.xpath('.//m:oMath|.//m:oMathPara'),i
        old=para.text
        if para.runs:
            para.runs[0].text=text
            for run in para.runs[1:]:run.text=''
        else:para.add_run(text)
        changes.append({'original_paragraph':i,'before':old,'after':text})

    def insert_before(anchor,text,style='Normal',bold=False):
        para=anchor.insert_paragraph_before(text,style)
        if style=='Normal' and body_pr is not None:
            if para._p.pPr is not None:para._p.remove(para._p.pPr)
            para._p.insert(0,deepcopy(body_pr))
        if bold:
            for run in para.runs:run.bold=True
        return para

    def add_table(anchor,heads,values,widths):
        t=doc.add_table(rows=1,cols=len(heads));t.style=doc.tables[0].style
        t.autofit=False
        for col,width in zip(t.columns,widths):col.width=Cm(width)
        for cell,text in zip(t.rows[0].cells,heads):cell.text=text
        for value in values:
            for cell,text in zip(t.add_row().cells,value):cell.text=str(text)
        for row_i,row in enumerate(t.rows):
            tr=row._tr.get_or_add_trPr()
            no_split=OxmlElement('w:cantSplit');tr.append(no_split)
            if row_i==0:
                repeat=OxmlElement('w:tblHeader');tr.append(repeat)
            for cell,width in zip(row.cells,widths):
                cell.width=Cm(width)
                if row_i==0:
                    fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'EAF1F6');cell._tc.get_or_add_tcPr().append(fill)
                for para in cell.paragraphs:
                    para.paragraph_format.line_spacing=1.15;para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(3)
                    para.paragraph_format.first_line_indent=Cm(0)
                    for run in para.runs:
                        run.font.size=Pt(10);run.bold=row_i==0
                        run.font.name='Times New Roman'
                        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
        anchor._p.addprevious(t._tbl)
        return t

    def image(anchor,name,caption):
        para=insert_before(anchor,'');para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent=Cm(0);para.paragraph_format.keep_with_next=True
        para.add_run().add_picture(str(HERE/'assets'/(name+'.png')),width=Cm(14.8))
        insert_before(anchor,caption,'题注1')

    replace(21,'OpenACM-Assistant是面向数字存内计算的本地智能电路设计平台。作品以“连续对话、目标驱动的自由探索、真实工具验证”为核心：用户给出电路规格、精度预算、优化目标与运行预算，系统在合法设计空间中逐轮选择候选、执行评价并利用反馈继续探索，最终输出可复查的可行方案、权衡关系与实验报告。')
    replace(22,'本项目自研连续自然语言交互、探索调度、任务级脚本编排、独立经验库和证据审计模块，复用OpenACM/OpenACMv2的电路生成、PEA-GNN及ACCO相关基础。现有能力覆盖自由探索、架构—器件分层评价、同条件方法对照和跨模型经验复用；已接入50000组公开8位乘法器参考配置，并完成真实乘法器探索、SRAM多PVT评价及16×8参考DCiM后端验证。')
    replace(37,'平台建立连续自然语言设计入口，使用户能够直接描述电路功能、误差预算、PVT条件和优化目标。对于探索任务，用户无需枚举每一个参数组合；系统构造合法候选、保留固定规格，在候选数与墙钟预算内组织多轮选择和验证。执行反馈、失败记录与当前最佳结果共同成为下一轮决策依据，历史任务通过本地经验库积累为可复用资料。')
    replace(38,'系统复用PEA-GNN对原生8/16位近似乘法器候选进行误差与PPA预测；SRAM当前以阵列SPICE及显式bank外围、面积估算完成评价，并保留TabPFN等表格代理模型的后续适配方向。在此基础上组织架构搜索、入选项复核、固定架构器件调整与指定PVT验证；参考PE任务接入OpenROAD等工具，输出RTL/网表、ODB、DEF、SPEF、SDF与报告等实际产物。')
    replace(49,'系统需要识别位宽、容量、供电、温度、工艺角、器件尺寸、误差阈值和优化目标，并将其转换为结构化参数。能力咨询、参数解释、运行、结果分析与后续修改共享设计上下文；用户既可指定单点配置，也可只给出目标、约束和预算，由系统组织合法空间中的自由探索。')
    replace(50,'探索过程中需要分别管理基线、固定条件、可调变量、候选预算、代理查询及真实EDA复核，保存每轮候选菜单、模型选择和工具反馈。系统应提供已评价可行最优、Pareto候选和停止原因，并将参数、指标、失败与环境统一归档，支持后续比较、复测和有条件的继续探索。')
    replace(57,'本地大模型承担自然语言理解、结构化规划与候选选择；自研外部执行器负责规格校验、预算调度、工作目录和证据归档。底层复用OpenACM的电路生成与仿真组件、原生PEA-GNN及NSGA-II等算法代码，并参考ACCO组织架构—器件两级流程。各模块通过受约束的参数接口协作，新的配置与辅助文件由执行器生成，历史观测写入项目独立经验库。')
    replace(58,'平台将语言决策与数值评价分工：语言模型根据目标、近期历史和允许菜单选择下一项实验，PEA-GNN在支持的乘法器范围内提供快速预测，RTL、SPICE、综合与STA提供相应复核。SRAM当前采用阵列SPICE与bank估算；TabPFN属于可扩展的表格代理方向，不作为本报告已验证的默认执行环节。')
    replace(60,'OpenACM/OpenACMv2提供电路编译、代理评价和协同优化的基础组件。本作品在原始程序只读的条件下，通过普通工作副本和外部适配层组织这些能力，增加自然语言任务契约、自由探索调度、实测验收及跨任务记录。底层组件的复用与上层新增功能在接口及证据中分别说明。')
    replace(62,'SRAM评价使用固定容量与输出字宽约束构建合法组织，调用实际阵列SPICE获得读写延迟、功耗和噪声容限，再结合bank片选、复用与面积模型进行组织比较。原框架中的TabPFN思路可用于后续表格代理适配；本报告所列SRAM实验的依据为SPICE与显式模型估算，未以TabPFN预测替代实测。')
    # Keep all three editable Office equations while correcting their implementation scope.
    para=p[63];maths=para._p.xpath('./m:oMath');assert len(maths)==3
    for child in list(para._p):
        if child.tag not in (qn('w:pPr'),qn('m:oMath')):para._p.remove(child)
    def run_at_math(text,before=None):
        run=para.add_run(text)
        if before is not None:before.addprevious(run._r)
    run_at_math('面向应用级协同设计，可将乘法器与SRAM的目标统一表述为：联合设计变量为 ',maths[0])
    run_at_math('，优化目标 ',maths[1]);run_at_math('，满足精度约束 ',maths[2])
    run_at_math('。当前平台分别组织乘法器与SRAM两类探索，在各自任务内先检查硬约束，再比较PPA与Pareto关系；跨电路配对到同一PE的联合目标评价作为后续扩展，不将两类独立任务相加视为已完成联合寻优。')
    replace(64,'平台按ACCO思路组织架构选点、入选实测和固定架构的器件复验，支持LLM、随机与NSGA-II等已适配策略。参考PE后端以单独任务执行，器件级改动的收益需完成单元重新表征和系统回灌后再评价。阶段参数、父子关系、调用计数与实际结果均回传上层并归档。')
    replace(65,'本项目的架构搜索对照模块固定共同候选池、初始配置和物理评价条件，并将搜索种子与算术测试种子分开。在8位案例中，共同池含256项，每轮最多评价32个唯一架构、最多3项进入RTL/综合/STA复核；LLM、随机和NSGA-II在同一规则下比较。该模块属于助手新增的实验编排能力。[9，E2]')
    replace(67,'本项目采用本地部署的DeepSeek系列模型作为上层设计智能体，通过本地推理服务接入。模型接收用户需求、结构化设计上下文、适用的经验片段和工具反馈，输出任务操作、参数建议、合法菜单编号及结果解释；执行器据此调用底层工具，并按模板生成任务所需的配置与辅助文件。')
    replace(68,'本地部署使电路源码、网表和实验记录可以保留在本机。模型的中文理解和代码理解能力用于需求映射、流程规划、候选选择与反馈解释；实际Python、Shell、Tcl、SPICE测试文件由受控执行器依照已校验字段生成。模型接口和经验存储相互独立，符合结构化输出要求的其他本地模型可以复用同一工作流。')
    replace(70,'平台将能力咨询、参数解释、设计修改、自由探索、结果分析和下一轮实验组织为连续设计对话。用户可以先咨询指标，再提出目标与预算，随后围绕已测配置调整条件；明确字段写入结构化状态，任务结果及下一步建议直接返回终端。当前入口仍需通过原话多轮验收持续改善，示例通过不等于任意表达都已覆盖。')
    replace(73,'系统以SQLite索引和标准JSON/JSONL建立本地持久经验库，保存完整配置、测试范围、指标、失败分类、工具指纹与证据位置。检索区分同评价条件和仅供参考的历史；适配的规划及局部探索入口可将同范围历史配置加入候选并重新实测，避免将旧数值直接当作新条件下的验收结果。')
    replace(75,'公开配置与标签保存在datasets，实际运行观测保存在experience，两者分别标记来源。经验跨会话保留，标准JSON/JSONL支持导出、导入和模型更换；积累记录与检索增强不自动修改模型权重。正式架构方法对照默认不注入持久经验，经验开关的配对消融仍需独立实现与验证。')
    replace(80,'OpenACM-Assistant的特色能力包括目标驱动的自由探索、连续设计对话、架构—器件两级任务编排、可替换的候选选择策略、跨模型持久经验，以及原始程序只读的任务级生成与执行。用户提交需求后，平台按预算推进候选评价，并把可行配置、Pareto关系、未达标项和工具失败汇总为可追溯结果。')
    replace(85,'OpenACM-Assistant划分为自然语言交互、AI规划调度、任务级轻量脚本、EDA执行、经验与审计五层。交互、预算控制、自由探索闭环及本地经验管理由本项目扩展；电路生成与部分算法组件复用OpenACM/OpenACMv2，EDA工具通过外部执行器调用。PEA-GNN已适配原生8/16位乘法器，SRAM目前采用SPICE与bank估算路径。')
    replace(89,'AI规划调度层维护目标、硬约束、可调域与固定条件，协调本地模型的下一候选选择和各阶段预算。代理预测、实际工具评价及多目标排序通过受控接口完成；模型不会因提出新候选而获得修改原始工程或执行任意程序的能力。')
    replace(95,'参数校验后，系统按适用范围检索历史记录并在独立目录生成任务文件。探索由外部执行器组织：构造合法且未评价的候选菜单，向模型提供当前已评价历史、失败与目标，再按选择调用对应工具。局部路径逐项实际验证，ACCO路径则将架构预测、入选复核和器件评价分开调度。')
    replace(98,'创新点一  连续对话驱动的自由探索')
    replace(99,'用户给出目标、约束和预算后，无需逐一填写所有候选配置。系统在允许空间中建立基线，逐轮提供合法菜单，由本地模型结合已评价结果选择下一项；工具反馈回到同一对话与搜索记录，形成“提出需求—选择候选—执行评价—利用反馈—继续探索”的闭环。')
    insert_before(p[100],'这种自主性体现在下一项实验的选择和预算内的迭代推进。固定规格与测试条件始终作为边界，模型选择记录和执行结果分别保存；范围外的拓扑修改需另行扩展接口，不能由一句“自由探索”绕过工程约束。')
    replace(101,'系统围绕每个任务生成所需的最小文件集合：乘法器任务组织RTL生成、固定测试向量与误差统计；SRAM任务组织阵列配置、器件参数与SPICE测试平台；完整后端任务进一步生成综合、布局布线与时序约束文件。所有修改在专属工作目录中完成，原始OpenACM内置程序保持只读。')
    replace(103,'创新点三  跨会话、跨模型的持久经验复用')
    replace(104,'经验库把配置、测量条件、结果、失败原因与证据位置保存为模型无关记录，支持跨会话检索与JSONL交换。公开参考和本地实测分开管理，历史方案需符合当前搜索域并重新验证；通过外部记忆积累设计经验，为更换模型或后续代理训练准备数据，而不将其表述为自动更新权重。')
    replace(106,'在ACCO两级思想上增加外部任务编排：架构阶段探索压缩器分配或SRAM组织，入选项完成相应复核后，再在固定架构下调整器件。每一级保存父子关系、基线和评价来源，便于从最终候选追溯其产生过程。')
    replace(107,'乘法器架构域包含允许的压缩器类型与逐位置分配；SRAM架构域包含行数、列数、子阵列数量及已支持的列复用比，并满足固定容量与输出接口。原生8/16位近似乘法器可用PEA-GNN筛选，SRAM当前用阵列SPICE加bank估算评价，入选后按指定条件开展器件与PVT任务。')
    replace(109,'创新点五  可验证的搜索过程与统一方法对照')
    replace(110,'平台区分预测、模型估算和真实工具测量，将每轮菜单、选择、失败、约束判定及停止原因归档。方法对照使用共同池、统一基线、分离种子和相同复核预算；成功指标与失败分母一同报告，使自由探索的过程和最终结果都能被检查，而不以模型的文字判断代替验收。')

    # New detailed feature section follows the original first figure.
    anchor=p[113]
    insert_before(anchor,'（四）目标驱动的自由探索工作流','Heading 2')
    insert_before(anchor,'自由探索是平台的主要交互能力之一：用户主要指定“要优化什么、必须满足什么、哪些条件固定、最多运行多久”，系统在已实现的设计域内安排下一项实验。以8位近似乘法器为例，可开放统一压缩器或逐位置分配，以MRED为约束，在综合面积或PDP目标下进行有限预算搜索。')
    insert_before(anchor,'闭环从同条件基线开始。执行器构造合法候选，局部探索菜单覆盖模板类型、当前最佳附近的单参数变化和多参数组合；ACCO阶段从明确的架构或器件域生成候选。LLM读取近期已评价记录和失败，在允许菜单中选择，程序再完成范围校验与工具调用。新的实测或预测按照来源进入下一轮历史，不能提前宣布尚未测量的候选达标。')
    insert_before(anchor,'策略可按阶段选择LLM、随机或NSGA-II。用户既可进行单点仿真，也可逐轮局部探索，或选择先架构后器件的研究流程；正式方法对照使用独立入口，确保共同候选池和评价预算一致。系统据配置维护可行最优与Pareto候选，达到候选数、墙钟时限、空间耗尽或中断条件后停止并说明原因。')
    image(anchor,'free_exploration_loop','图X  目标驱动的自由探索与工具反馈闭环')
    insert_before(anchor,'表X  特色功能、交互方式与交付结果','题注1')
    add_table(anchor,['特色功能','系统承担的工作','固定边界与输出'],[
        ['连续设计对话','将咨询、修改、执行和结果分析关联到同一设计','保留结构化条件；输出计划、指标及后续建议'],
        ['目标驱动自由探索','构造合法菜单，结合反馈选择下一候选并继续评价','固定规格/测试；输出可行最优、Pareto及停止原因'],
        ['两级任务编排','架构选点、入选复核、固定架构的器件/PVT评价','保存阶段父子关系；区分预测、SPICE和估算'],
        ['持久经验与公开参考','按测试条件检索配置与失败记录，支持JSONL交换','参考与实测分开；历史候选仍重新验收'],
        ['统一方法对照','在同池、同基线、同评价条件下调度多种策略','分离搜索/测试种子；记录调用次数与墙钟成本'],
        ['任务隔离与追溯','生成轻量脚本，保存配置、网表、日志与哈希','原始程序只读；所有新文件进入专属任务目录'],
    ],[3.0,5.5,6.3])
    insert_before(anchor,'后续复验可引用保存的完整配置：SRAM使用best_experiment保留总容量、组织与PVT，乘法器保留压缩器位置分配和测试条件。旧单级探索支持检查点续跑；完整ACCO当前用原配置新建任务复验，尚未开放中途恢复。这样可以持续推进研究，同时保留每次尝试的文件与判定。')
    replace(113,'（五）安全与可复现机制')

    # Correct an unsupported attribution in the tools table.
    doc.tables[0].rows[3].cells[0].text='TabPFN（扩展方向）'
    doc.tables[0].rows[3].cells[1].text='SRAM表格代理预测'
    doc.tables[0].rows[3].cells[2].text='后续适配；本次SRAM采用SPICE与bank估算'
    doc.tables[0].rows[4].cells[0].text='LLM／随机／NSGA-II'
    doc.tables[0].rows[4].cells[2].text='按阶段选择策略并进行公平方法对照'
    doc.tables[2].rows[2].cells[1].text='目标、约束、预算及候选反馈调度'
    doc.tables[2].rows[4].cells[1].text='电路生成、代理接口与实际EDA评价'
    # Add the documented autonomous-search evidence to representative results.
    cells=doc.tables[1].add_row().cells
    for cell,text in zip(cells,['自由探索闭环','128条候选记录；18项满足约束','8位近似乘法器；127次模型选点；面积目标']):cell.text=text
    replace(129,'图3依据同一任务的route.def生成，展示单元锚点与布线坐标。对应任务保存ODB、DEF、SPEF、SDF和分析报告；该图为DEF坐标视图，参考流程未承诺自动导出GDSII，也不将详细布线DRC=0等同完整工艺DRC/LVS签核。[9，E6]')

    anchor=p[135]
    insert_before(anchor,'（四）自由探索闭环的实际验证','Heading 2')
    insert_before(anchor,'历史任务'+TASK+'展示了真正的多轮候选选择与实际验证。任务固定8位近似乘法器、seed=42、4096组测试、Nangate45 typical综合与MRED≤0.02要求，目标优先最小化综合面积、再比较MRED，预算128条候选记录。第一项为基线，其余127项由本地DeepSeek从合法菜单选择，并逐项执行RTL、综合及门级对照。[10]')
    insert_before(anchor,'任务保存18条满足约束、86条误差未达标和24条执行失败记录，形成6个Pareto候选。最佳已测项为第6号AKbar2配置，MRED=0.0136390362，综合面积291.802μm²；相对本任务ACCI1基线322.924μm²减少约9.64%。这一结果体现了精度预算与面积目标的取舍，不表示误差也同时降低。[10]')
    insert_before(anchor,'表X  自由探索任务的真实执行记录','题注1')
    add_table(anchor,['项目','本次记录','评价含义'],[
        ['候选及来源','128条；1基线＋127次DeepSeek选点','每条都有相应候选与选择记录，不写成128次成功'],
        ['状态分布','18满足约束／86未达标／24执行失败','保留全部失败与分母，不以缺测值充当0'],
        ['最佳已测配置','第6号；统一AKbar2','仅对本任务有限候选成立'],
        ['MRED','0.0136390362≤0.02','4096组固定测试通过，不是穷举证明'],
        ['综合面积','322.924→291.802μm²，约−9.64%','与同任务基线比较；不是布线面积或LLM策略优势'],
        ['Pareto与停止','6项；达到128条候选预算后停止','保留权衡方案，不宣称全空间最优'],
        ['源工程校验','2631文件；unchanged=true','原始OpenACM未改动'],
    ],[3.2,5.7,5.9])
    image(anchor,'free_exploration_results','图X  自由探索的可行面积轨迹与实测误差—面积分布（24条执行失败无数值点）')
    insert_before(anchor,'曲线在第6条后未出现更小的可行面积，表明后续预算并未继续改善该目标。这份记录验证了自动选点、反馈和归档链路，不能单独证明LLM搜索优于随机或NSGA-II。独立的三方法同池实验已有15轮尝试、14轮完整完成，本组未显示LLM更低的中位PDP；算法优劣按该对照口径判断。[9，E2]')
    insert_before(anchor,'另一个任务'+ACCO+'完成LLM架构选点、2次GNN查询、1次RTL/综合/STA与1项器件评价，说明上述语言选点已接入两级流程。它是小预算链路验证，与128条逐项实测的旧单级任务分别记录，不把两者混写成同一种搜索规模。[10]')
    replace(135,'（五）设计数据与成果管理')
    replace(139,'（六）应用场景与科研价值')
    replace(148,'本项目在OpenACM/OpenACMv2基础上，将本地模型、确定性执行器与EDA工具连接为连续设计平台。作品特色在于以自然语言给定目标和边界后，能够在允许设计域内逐轮选点、利用工具反馈推进探索；同时以独立任务目录、阶段记录及持久经验将设计过程保存为可复查、可复用的研究资产。')
    replace(149,'现有证据包括128条候选的模型选点与实测闭环、50000组公开参考接入、64bit SRAM的9个器件/PVT子任务，以及16×8精确DCiM的65536组乘法输入验证和12个后端阶段。自由探索任务在MRED≤0.02下得到面积较本任务基线小9.64%的已测配置；参考DCiM满足本次时序与布线验收。以上证明工程链路及可追溯性，策略优越性与应用级收益仍需独立对照。[9、10]')
    replace(151,'后续将完善自由探索的候选质量、停止策略和多轮原话理解，补齐独立代理验证与经验开关消融；在严格区分预测和实测的基础上，引入适合的表格代理与不确定性评价。进一步扩展应用级误差、更多工艺库及电路类型，按统一测试范围积累训练与比较数据。')
    replace(158,'参考DCiM配置为16行×8位、exact、backend=asic、时钟10ns；SRAM配置为64bit、输出4bit、3种PD×3组PVT；自由探索案例固定8位、seed=42、4096组和MRED≤0.02。旧单级记录、ACCO两级验证与正式方法对照各用独立编号，不将其预算或评价来源混计。')
    replace(160,'软件构成包括连续对话Schema、本地模型接口、探索候选与预算调度、任务级脚本生成、OpenACM只读适配、PEA-GNN接口、SQLite经验索引、标准JSONL交换，以及RTL/SPICE/STA/OpenROAD后端。模型从合法菜单输出选择，执行器负责生成文件与运行工具；各部分可以分别替换和验证。')
    for text in [
        ['自由探索：128条候选与127次模型选点',TASK],
        ['LLM选点接入两级探索',ACCO],
        ['三方法同池架构搜索对照','20261027_204422_search_comparison_d356cf'],
    ]:
        cells=tables[6].add_row().cells
        for cell,value in zip(cells,text):cell.text=value
    insert_before(p[171],'[10] OpenACMAssistant本机探索记录：'+TASK+'；'+ACCO+'。本次2026年10月29日从result.json、逐轮决策及完整性文件重新核对。')

    # Renumber captions in their actual document order, retaining existing figures.
    numbers={'图':0,'表':0}
    for para in doc.paragraphs:
        if para.style.name!='题注1':continue
        match=re.match(r'^(图|表)(?:\d+|X)\s+(.*)',para.text)
        if match:
            key=match[1];numbers[key]+=1
            para.text=f'{key}{numbers[key]}  {match[2]}'
    # Fix excessive fallback-font line heights when opening the WPS source in
    # another office suite; 18pt retains the intended 12pt x 1.5 body spacing.
    for element in doc._element.xpath('.//w:p'):
        props=element.get_or_add_pPr()
        for old in props.findall(qn('w:snapToGrid')):props.remove(old)
        snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');props.append(snap)
    for index,para in enumerate(doc.paragraphs):
        if index < 20:
            if not para.text and not para._p.xpath('.//w:drawing|.//w:pict'):
                para.paragraph_format.line_spacing=Pt(12)
                para.paragraph_format.space_before=Pt(0)
                para.paragraph_format.space_after=Pt(0)
            elif para.text:
                size=max([r.font.size.pt for r in para.runs if r.font.size] or [16])
                para.paragraph_format.line_spacing=Pt(size*1.4)
            continue
        if para.style.name=='Normal' and para.text.strip():
            para.paragraph_format.line_spacing=Pt(18)
        elif para.style.name in ('Heading 1','Heading 2'):
            para.paragraph_format.keep_with_next=True
            para.paragraph_format.line_spacing=Pt(22 if para.style.name=='Heading 1' else 20)
        elif para.style.name=='题注1':
            para.paragraph_format.line_spacing=Pt(14)
        elif not para.text and para._p.xpath('./w:pPr/w:sectPr'):
            para.paragraph_format.line_spacing=Pt(1)
            para.paragraph_format.space_before=Pt(0)
            para.paragraph_format.space_after=Pt(0)
    output=HERE/(STEM+'.docx');doc.save(output)
    assert sha(SOURCE)==original_hash,'Original user report changed'
    # Equations, source images, sections and original non-target content stay in the copy.
    original=Document(SOURCE);updated=Document(output)
    assert len(original.sections)==len(updated.sections)==3
    assert len(updated.inline_shapes)==4 and len(updated.tables)==9
    assert len(original._element.xpath('.//m:oMath'))==len(updated._element.xpath('.//m:oMath'))
    for relation in original.part.rels.values():
        if relation.reltype.endswith('/image'):
            assert any(r.reltype.endswith('/image') and r.target_part.blob==relation.target_part.blob for r in updated.part.rels.values())
    all_text='\n'.join(v.text for v in updated.paragraphs)
    for required in ['自由探索','127','9.64%','不自动修改模型权重','尚未开放中途恢复','未承诺自动导出GDSII']:
        assert required in all_text,required
    write_json(HERE/'evidence/revision_log.json',{'date':'2026-10-29','source':str(SOURCE),
        'source_sha256':original_hash,'output':str(output),'output_sha256':sha(output),
        'changed_paragraphs':changes,'captions':numbers,'sections_preserved':3,'original_images_preserved':2,
        'equations_preserved':len(updated._element.xpath('.//m:oMath')),
        'paragraphs':len(updated.paragraphs),'tables':len(updated.tables),'images':len(updated.inline_shapes)})
    md=['# 技术报告：特色功能完善版','','正文保留原稿结构；公式和封面以Word为准。','']
    for child in updated._element.body:
        if child.tag==qn('w:p'):
            para=Paragraph(child,updated._body)
            if not para.text.strip():continue
            prefix='## ' if para.style.name=='Heading 1' else '### ' if para.style.name=='Heading 2' else ''
            md.extend([prefix+para.text,''])
        elif child.tag==qn('w:tbl'):
            from docx.table import Table
            table=Table(child,updated._body);values=[[c.text.replace('\n','<br>') for c in r.cells] for r in table.rows]
            md.extend(['| '+' | '.join(values[0])+' |','| '+' | '.join(['---']*len(values[0]))+' |'])
            md.extend('| '+' | '.join(row)+' |' for row in values[1:]);md.append('')
    (HERE/(STEM+'.md')).write_text('\n'.join(md))
    print(output)


if __name__=='__main__':main()

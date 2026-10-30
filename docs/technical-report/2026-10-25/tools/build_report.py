"""Build a self-contained report from existing evidence; never run EDA or modify sources."""
from __future__ import annotations

import argparse
import csv
from datetime import datetime
import hashlib
import json
from pathlib import Path
import re
import subprocess
import sys
from zoneinfo import ZoneInfo

HERE = Path(__file__).resolve().parents[1]
ROOT = HERE.parents[2]
STEM = "OpenACMAssistant_技术报告"
BASELINE = "8566e1ac9eff28dc668404e10480ef2ac664a70b"
RUNS = {
    "E01": "20261025_134117_acco_sram_fb675e",
    "E02": "20261025_132043_acco_multiplier_052c4b",
    "E03": "20261025_132307_acco_multiplier_4504d7",
    "E04": "20261025_132333_acco_multiplier_5f789e",
    "E05": "20261025_134246_acco_multiplier_b21791",
    "E06": "20261025_133552_acco_multiplier_84cbab",
    "E07": "20261025_132614_acco_multiplier_ae5f8b",
}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read(path):
    return json.loads(Path(path).read_text())


def write(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n")


def collect():
    evidence = {"schema": "openacm-report-evidence/v1", "report_date": "2026-10-25",
        "collected_at": datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(),
        "code_commit": BASELINE,
        "source_checkout_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(),
        "scope": "Existing local simulation and audit records; no new EDA or model inference", "runs": {}}
    for identity, name in RUNS.items():
        task = ROOT / "runs" / name
        result = read(task / "result.json")
        integrity = read(task / "source_integrity.json")
        assert integrity["unchanged"] and integrity["checked_files"] == 2631, name
        assert result["status"] == ("infeasible" if identity in ("E06", "E07") else "completed"), name
        stages = {k: {a: b for a, b in v.items() if a != "records"} for k, v in result["stages"].items()}
        for k, stage in stages.items():
            stage["records"] = [{a: b for a, b in r.items() if a in ("candidate", "parent", "parameters", "status", "metrics", "constraints", "evidence_kind", "proposed_by", "task", "elapsed_s")}
                for r in result["stages"][k]["records"]]
        evidence["runs"][identity] = {"task": "runs/" + name, "result_sha256": sha(task / "result.json"),
            "config": result["config"], "status": result["status"], "counters": result["counters"],
            "best_design": result.get("best_design"), "best_device_designs": result.get("best_device_designs", []),
            "stages": stages, "source_integrity": integrity,
            "pvt_coverage": [read(p) for p in sorted(task.glob("stages/*/*/coverage.json"))]}
    assert any(r.get("proposed_by") == "deepseek" for r in evidence["runs"]["E04"]["stages"]["architecture"]["records"])
    bad = evidence["runs"]["E06"]
    assert bad["best_design"] is None and bad["counters"]["device_candidates"] == 0
    conditions = bad["stages"]["validation"]["records"][0]["constraints"]
    assert next(c["passed"] for c in conditions if c["metric"] == "mred")
    assert not next(c["passed"] for c in conditions if c["metric"] == "pdp_fj")
    preview = ROOT / "runs/20261025_133919_planning_6cad2c/plan.json"
    evidence["E08"] = {"path": str(preview.relative_to(ROOT)), "sha256": sha(preview), "plan": read(preview)}
    assert evidence["E08"]["plan"]["constraints"] == []
    summary = sorted((ROOT / "runs/20261025_134015_acco_multiplier_eff2af/reports/summaries").glob("*/metadata.json"))[-1]
    evidence["E09"] = {"path": str(summary.relative_to(ROOT)), "sha256": sha(summary), "metadata": read(summary)}
    assert evidence["E09"]["metadata"]["model_used"]
    history = ROOT / "runs/20261010_144501_full_flow_89aabc"
    evidence["E10"] = {p.parent.name: {"sha256": sha(p), "result": read(p)} for p in history.glob("tasks/*/result.json")}
    evidence["E11"] = read(ROOT / "verification/release_regression.json")
    # The temporary release directory is not part of the portable report.
    evidence["E11"].pop("snapshot", None)
    assert evidence["E11"]["status"] == "passed" and all(c["exit_code"] == 0 for c in evidence["E11"]["checks"])
    evidence["source_documents"] = {name: sha(ROOT / name) for name in ["app/acco.py", "app/acco_models.py", "app/acco_physics.py",
        "app/acco_spec.py", "app/worker.py", "app/guard.py", "app/conversation.py", "app/experience.py", "app/result_summary.py",
        "docs/验证记录.md", "docs/公开数据接入.md", "docs/持久经验库.md"]}
    evidence["paper_sha256"] = {name: sha(ROOT.parent / name) for name in ("DATE.pdf", "DAC.pdf")}
    for name, digest in evidence["source_documents"].items():
        versioned = subprocess.check_output(["git", "show", BASELINE + ":" + name], cwd=ROOT)
        assert hashlib.sha256(versioned).hexdigest() == digest, "Review changed source: " + name
    write(HERE / "evidence_snapshot.json", evidence)
    return evidence


def figures(evidence):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.font_manager import FontProperties
    from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
    font = FontProperties(fname="/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc")
    plt.rcParams.update({"font.family": font.get_name(), "svg.fonttype": "path", "axes.unicode_minus": False,
        "font.size": 10, "figure.facecolor": "white", "savefig.facecolor": "white"})
    navy, teal, gray = "#15324F", "#087F83", "#65778A"
    def canvas(height=4.5):
        fig, ax = plt.subplots(figsize=(7.2, height))
        ax.set(xlim=(0, 1), ylim=(0, 1)); ax.axis("off")
        fig.subplots_adjust(left=.015, right=.985, top=.98, bottom=.025)
        return fig, ax
    def box(ax, x, y, w, h, text, fill="#EAF3F7", edge=navy, size=10, color=navy):
        ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.008,rounding_size=0.015",
            facecolor=fill, edgecolor=edge, linewidth=1.2))
        ax.text(x+w/2, y+h/2, text, ha="center", va="center", fontsize=size, color=color, fontproperties=font, linespacing=1.5)
    def arrow(ax, start, end, dashed=False):
        ax.add_patch(FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=12, linewidth=1.2,
            color=gray, linestyle="--" if dashed else "-", connectionstyle="arc3"))
    def save(fig, name):
        for ext in ("png", "svg"):
            fig.savefig(HERE / "assets" / (name + "." + ext), dpi=240)
        plt.close(fig)

    fig, ax = canvas(4.0)
    box(ax, .08, .79, .84, .15, "OpenACMAssistant 上层助手\n连续对话 · 实验规格 · 搜索调度 · 证据总结", "#15324F", color="white", size=11)
    for x, text in [(.03, "乘法器编译器\n精确 / 近似 / Log\n位宽与压缩器分配"), (.35, "SRAM 电路内核\n6T 与外围配置\nSPICE / 读写 / SNM"), (.67, "SRAM 与 PE 编译器\n行为存储 / 控制 / 集成\nRTL 与参考宏接口")]:
        box(ax, x, .43, .29, .23, text)
        arrow(ax, (x+.145, .78), (x+.145, .68))
        arrow(ax, (x+.145, .42), (x+.145, .29))
    box(ax, .03, .10, .93, .18, "EDA 工具执行与结果\nXyce · Icarus/vvp · Yosys · OpenSTA · OpenROAD · VCS", "#E7F5F0", teal, size=10)
    ax.text(.5, .025, "原始内核只读；生成与仿真在任务普通副本中执行", ha="center", color=gray, fontsize=9, fontproperties=font)
    save(fig, "kernel_stack")

    fig, ax = canvas(4.1)
    for x, title in [(.02, "中文需求 / 连续追问\n设计、咨询与澄清"), (.365, "本地 DeepSeek\n结构化计划与决策"), (.71, "规格校验 / ACCO\n约束、空间与预算")]:
        box(ax, x, .72, .27, .20, title)
    arrow(ax, (.30, .82), (.355, .82)); arrow(ax, (.645, .82), (.70, .82))
    for x, title in [(.02, "持久经验 / 公开参考\n条件匹配、版本和来源"), (.365, "实测验收 / 中文总结\n事实渲染与完整审计"), (.71, "隔离执行器 / 内核\nGNN、SPICE、RTL、EDA")]:
        box(ax, x, .36, .27, .22, title, "#E7F5F0", teal)
    arrow(ax, (.845, .71), (.845, .59)); arrow(ax, (.70, .47), (.645, .47)); arrow(ax, (.355, .47), (.30, .47))
    arrow(ax, (.155, .59), (.155, .71), True)
    box(ax, .08, .06, .84, .15, "全过程规则：原始程序只读 · 预测/估算/实测分离\n固定测试条件 · 预算停止 · 参数、日志与SHA可追溯", "#F4F5F8", "#B2BBC4", size=9.5)
    save(fig, "system_architecture")

    fig, ax = canvas(5.2)
    ax.text(.255, .96, "近似乘法器", ha="center", color=navy, fontsize=12, fontproperties=font, weight="bold")
    ax.text(.745, .96, "SRAM bank", ha="center", color=navy, fontsize=12, fontproperties=font, weight="bold")
    left = ["固定原生8/16位结构\n逐位置压缩器组合", "PEA-GNN 架构预测\nLLM / 随机 / NSGA-II", "RTL → 综合 → STA活动\n按真实误差与PPA验收", "固定逻辑，分组宽度优化\nSPICE PPA + 16组目标真值"]
    right = ["明确总容量与输出字宽\nr × c × 阵列数 = 容量", "固定标称器件，搜索组织\n阵列SPICE + bank估算", "入选bank，固定组织\n保持容量、复用比与接口", "器件尺寸 + 指定PVT/MC\n读写、SNM与最坏条件验收"]
    for x, labels in ((.035, left), (.525, right)):
        for i, text in enumerate(labels):
            y = .75 - i*.20
            box(ax, x, y, .44, .15, text, "#EAF3F7" if i < 2 else "#E7F5F0", navy if i < 2 else teal)
            if i: arrow(ax, (x+.22, y+.195), (x+.22, y+.157))
    ax.text(.5, .025, "统一输出：阶段父子关系、有效候选、验收依据、同条件baseline与复验配置", ha="center", color=gray, fontsize=9, fontproperties=font)
    save(fig, "two_level_flow")

    rows = evidence["runs"]["E01"]["stages"]["device_architecture_0002"]["records"]
    keys = ["min_snm_v", "bank_read_delay_ns_estimate", "bank_read_power_uw_estimate"]
    titles = ["最小 SNM / V", "读延迟估算 / ns", "读功耗估算 / µW"]
    fig, axes = plt.subplots(1, 3, figsize=(7.2, 2.8))
    for ax, key, title in zip(axes, keys, titles):
        values = [r["metrics"][key] for r in rows]
        ax.bar([0, 1], values, color=["#ABB8C3", teal if key != keys[-1] else "#C08A2D"], width=.58)
        ax.set(xticks=[0, 1], xticklabels=["PD135", "PD180"], title=title, ylim=(0, max(values)*1.25))
        for i, value in enumerate(values): ax.text(i, value+max(values)*.035, f"{value:.6g}", ha="center", fontsize=9)
        ax.spines[["right", "top"]].set_visible(False)
        ax.tick_params(labelsize=8)
        ax.grid(axis="y", alpha=.18); ax.set_axisbelow(True)
    fig.suptitle("固定8×4×2组织；TT与SS两角，各MC=2、相同种子", fontsize=10, y=.99)
    fig.tight_layout(rect=(0, 0, 1, .92))
    save(fig, "sram_tradeoff")
    with (HERE / "assets/sram_tradeoff.csv").open("w") as stream:
        writer = csv.writer(stream); writer.writerow(["PD_width_nm"] + keys)
        for r in rows: writer.writerow([r["parameters"]["devices"]["pd_width_nm"]] + [r["metrics"][k] for k in keys])


def word():
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Cm(2)
    section.header_distance = section.footer_distance = Cm(1)
    section.different_first_page_header_footer = True
    def configure(style, size, east="Noto Serif CJK SC", bold=False, color="243746"):
        style.font.name = "DejaVu Sans"; style.font.size = Pt(size); style.font.bold = bold
        style.font.color.rgb = RGBColor.from_string(color)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east)
    configure(doc.styles["Normal"], 10.5)
    normal = doc.styles["Normal"].paragraph_format
    normal.line_spacing = 1.28; normal.space_after = Pt(5); normal.widow_control = True
    for n, size in ((1, 16), (2, 12.5), (3, 11)):
        style = doc.styles[f"Heading {n}"]
        configure(style, size, "Noto Sans CJK SC", True, "15324F")
        style.paragraph_format.space_before = Pt(14 if n == 1 else 10)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    configure(doc.styles["Caption"], 9, "Noto Sans CJK SC", color="65778A")
    doc.styles["Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    configure(doc.styles["Header"], 8, "Noto Sans CJK SC", color="65778A")
    configure(doc.styles["Footer"], 8, "Noto Sans CJK SC", color="65778A")
    section.header.paragraphs[0].text = "OpenACMAssistant  |  本地智能设计与两级探索"
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("技术报告 V1.0  ·  ")
    for instr in ("PAGE", "NUMPAGES"):
        if instr == "NUMPAGES": footer.add_run(" / ")
        field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), instr); footer._p.append(field)
    doc.core_properties.title = "OpenACMAssistant 技术报告"
    doc.core_properties.subject = "近似存内计算电路的本地智能设计与ACCO两级探索"
    doc.core_properties.author = "OpenACMAssistant 项目"
    doc.core_properties.keywords = "OpenACM, DeepSeek, ACCO, GNN, SRAM, 设计空间探索"
    for text, size, color, before in [("OPENACM", 40, "15324F", 75), ("ASSISTANT", 22, "087F83", 0),
        ("参赛技术报告", 21, "15324F", 32), ("本地大模型驱动的近似存内计算\n电路设计与两级探索系统", 16, "243746", 16)]:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        run = p.add_run(text); run.bold = True; run.font.size = Pt(size); run.font.color.rgb = RGBColor.from_string(color)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Noto Sans CJK SC")
    p = doc.add_paragraph("V1.0  ·  2026年10月25日\n代码基准 8566e1a")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(44)
    p = doc.add_paragraph("语言驱动设计  /  两级协同探索  /  实测证据闭环")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(20)
    doc.add_page_break()
    def inline(paragraph, text):
        for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
            if not token: continue
            run = paragraph.add_run(token[2:-2] if token.startswith("**") else token[1:-1] if token.startswith("`") else token)
            if token.startswith("**"): run.bold = True
            if token.startswith("`"): run.font.name = "DejaVu Sans Mono"; run.font.size = Pt(9)
    lines = (HERE / (STEM + ".md")).read_text().splitlines()
    start = lines.index("## 摘要")
    i, code = start, False
    while i < len(lines):
        line = lines[i]
        if not line.strip(): i += 1; continue
        if line.startswith("```"):
            code = not code; i += 1; continue
        if code:
            p = doc.add_paragraph(); p.paragraph_format.line_spacing = 1.1; p.paragraph_format.space_after = Pt(1)
            run = p.add_run(line); run.font.name = "DejaVu Sans Mono"; run.font.size = Pt(8.1)
            shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "F1F4F7"); p._p.get_or_add_pPr().append(shade)
        elif line.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.*?)\)", line)
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(HERE / match[2]), width=Inches(6.6))
            doc.add_paragraph(match[1], "Caption")
        elif line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"\|[\s:|-]+\|$", lines[i]): table_lines.append(lines[i])
                i += 1
            cells = [[c.strip() for c in row.strip("|").split("|")] for row in table_lines]
            table = doc.add_table(rows=1, cols=len(cells[0])); table.style = "Table Grid"
            table.autofit = False
            count = len(cells[0])
            widths = ([1.6, 5.0] if count == 2 else [2.0, 1.1, 1.1, 1.1, 1.3] if count == 5 else
                [.45, 3.65, 2.5] if cells[0][0] == "编号" else [1.5, 2.55, 2.55])
            for column, width in zip(table.columns, widths): column.width = Inches(width)
            header = OxmlElement("w:tblHeader"); table.rows[0]._tr.get_or_add_trPr().append(header)
            for ri, row in enumerate(cells):
                target = table.rows[0] if ri == 0 else table.add_row()
                no_split = OxmlElement("w:cantSplit"); target._tr.get_or_add_trPr().append(no_split)
                for ci, value in enumerate(row):
                    target.cells[ci].width = Inches(widths[ci])
                    p = target.cells[ci].paragraphs[0]; inline(p, value)
                    p.paragraph_format.space_after = Pt(4); p.paragraph_format.space_before = Pt(3); p.paragraph_format.line_spacing = 1.12
                    if len(cells) <= 7 and ri < len(cells)-1: p.paragraph_format.keep_with_next = True
                    for run in p.runs:
                        run.font.size = Pt(8.5 if len(cells[0]) >= 5 else 9)
                        if ri == 0: run.bold = True; run.font.color.rgb = RGBColor.from_string("FFFFFF")
                    if ri == 0:
                        shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "15324F"); target.cells[ci]._tc.get_or_add_tcPr().append(shade)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        elif line.startswith("## "):
            title = line[3:]
            heading = doc.add_heading(title, 1)
            if re.match(r"(?:1 |2 |4 |7 |10 |12 )", title): heading.paragraph_format.page_break_before = True
        elif line.startswith("### "):
            heading = doc.add_heading(line[4:], 2)
            if line.startswith("### 12.2 "): heading.paragraph_format.page_break_before = True
        else:
            p = doc.add_paragraph(); inline(p, line)
        i += 1
    doc.save(HERE / (STEM + ".docx"))


def audit(check_pdf=False):
    from docx import Document
    source = (HERE / (STEM + ".md")).read_text()
    doc = Document(HERE / (STEM + ".docx"))
    assert len(doc.inline_shapes) == 4
    assert len(doc.tables) == 11, len(doc.tables)
    assert all((HERE / p).is_file() for p in re.findall(r"!\[[^]]*\]\(([^)]+)\)", source))
    snapshot = read(HERE / "evidence_snapshot.json")
    assert snapshot["code_commit"].startswith("8566e1a"), "Report baseline changed; review text and regenerate"
    rows = snapshot["runs"]["E01"]["stages"]["device_architecture_0002"]["records"]
    changes = {k: (rows[1]["metrics"][k]/rows[0]["metrics"][k]-1)*100 for k in ["min_snm_v", "bank_read_delay_ns_estimate", "bank_read_power_uw_estimate"]}
    for identity in ("E02", "E03"):
        metrics = snapshot["runs"][identity]["best_design"]["metrics"]
        for metric, precision in [("mred", 10), ("nmed", 10), ("synthesized_area_um2", 3), ("sta_delay_ns", 8), ("sta_dynamic_power_uw", 7), ("pdp_fj", 7)]:
            value = f"{metrics[metric]:.{precision}f}"
            assert value in source, (identity, metric, value)
    for value in ["648.757583", "0.1694452"]: assert value in source
    for value in changes.values(): assert f"{abs(value):.2f}%" in source
    manifest = {"report_baseline": snapshot["code_commit"], "generated_at": datetime.now(ZoneInfo("Asia/Shanghai")).isoformat(),
        "figures": len(doc.inline_shapes), "tables": len(doc.tables), "sram_relative_changes_percent": changes,
        "files": {str(p.relative_to(HERE)): {"sha256": sha(p), "bytes": p.stat().st_size} for p in HERE.rglob("*")
            if p.is_file() and "build" not in p.relative_to(HERE).parts and p.name != "manifest.json"},
        "checks": ["original task integrity", "expected completed/infeasible states", "real model choice and summary", "PDP constraint rejection",
            "reported values", "all embedded figures", "docx paragraph and table structure"], "status": "passed" if check_pdf else "awaiting_pdf_export_and_audit"}
    pdf = HERE / (STEM + ".pdf")
    if check_pdf:
        assert pdf.is_file(), "Export PDF before the final audit"
        text = subprocess.check_output(["pdftotext", "-layout", str(pdf), "-"], text=True)
        pages = [p for p in text.split("\f") if p.strip()]
        assert all(len(re.sub(r"\s+", "", page)) > 80 for page in pages), "Blank or nearly empty PDF page"
        assert "\ufffd" not in text
        compact = re.sub(r"\s+", "", text)
        for title in re.findall(r"^## ([0-9]+ .+)$", source, re.M):
            assert re.sub(r"\s+", "", title) in compact, "Missing PDF section: " + title
        for value in ["0.0136390362", "63.1574730", "0.1694452", "429.3007558"]: assert value in compact
        manifest["pdf_pages"] = len(pages)
        manifest["checks"] += ["PDF all chapters present", "PDF no empty pages or replacement glyphs", "PDF key numeric values"]
    write(HERE / "manifest.json", manifest)
    print("Report validated:", len(source), "characters,", len(doc.tables), "tables, 4 figures")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--audit-only", action="store_true")
    args = parser.parse_args()
    if not args.audit_only:
        figures(collect()); word()
    audit(check_pdf=args.audit_only)

"""Validate report data, actual replies, source checks and DOCX packaging."""
import hashlib
import json
from pathlib import Path
import zipfile

from docx import Document

HERE = Path(__file__).resolve().parents[1]
ROOT = HERE.parents[2]
e = json.loads((HERE / "evidence_snapshot.json").read_text())
word = HERE / "OpenACMAssistant_E1-E6问答与实验佐证报告.docx"
doc = Document(word)
paragraphs = [p.text for p in doc.paragraphs]
assert len(e["dialogue"]) == 12
assert all(t["metadata"].get("model_used") for t in e["dialogue"])
for i, turn in enumerate(e["dialogue"], 1):
    marker = "用户：" + turn["request"]
    begin = paragraphs.index(marker)
    end = next(n for n in range(begin+1, len(paragraphs)) if paragraphs[n].startswith("审计："))
    actual = paragraphs[begin+2:end]
    assert actual == turn["reply"].split("\n\n"), f"Reply changed at turn {i}"
assert len(e["new_jobs"]) == 3
assert all(j["integrity"]["unchanged"] for j in e["new_jobs"])
assert e["E1_repeat_check"]["rtl_sha256_equal"]
assert e["E1_repeat_check"]["first_test_is_prefix_of_second"]
assert e["E4_supplement"]["successful_observations"] == 44
assert e["E4_supplement"]["unique_designs"] == 17
assert e["E4_supplement"]["confusion_mred_002"]["pred_feasible_meas_infeasible"] == 1
assert e["tasks"]["E6"]["result"]["status"] == "completed"
assert len(e["E6_stages"]) == 12 and all(s["status"] == "completed" for s in e["E6_stages"])
assert len(doc.inline_shapes) == 7
for filename, expected in e["source_hashes"].items():
    assert hashlib.sha256((ROOT / filename).read_bytes()).hexdigest() == expected, filename
with zipfile.ZipFile(word) as archive:
    assert archive.testzip() is None
result = {
    "status": "passed", "dialogue_count": 12, "real_model_turns": 12,
    "new_jobs": [{"task": j["task"], "status": j["result"]["status"]} for j in e["new_jobs"]],
    "paragraphs": len(doc.paragraphs), "tables": len(doc.tables), "embedded_figures": len(doc.inline_shapes),
    "reply_verbatim": True, "source_hashes_passed": len(e["source_hashes"]),
    "docx_bytes": word.stat().st_size,
    "docx_sha256": hashlib.sha256(word.read_bytes()).hexdigest(),
    "pdf_render_verified": False,
}
(HERE / "validation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n")
print(json.dumps(result, ensure_ascii=False))
